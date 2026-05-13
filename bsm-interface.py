# ============================================================
# SISTEM KLASIFIKASI PENERIMA BSM - NAIVE BAYES
# Dashboard Machine Learning Modern Premium | Streamlit App
# Jalankan: streamlit run klasifikasi-bsm.py
# ============================================================

import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
from sklearn.preprocessing import StandardScaler, LabelEncoder
from sklearn.model_selection import train_test_split
from sklearn.naive_bayes import GaussianNB
from sklearn.metrics import accuracy_score, confusion_matrix
import io
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────
# KONFIGURASI HALAMAN AWAL
# ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Klasifikasi BSM | Naive Bayes",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─────────────────────────────────────────────────────────────
# CUSTOM CSS — NEO BRUTALISM + Y2K MODERN WEB (PERBAIKAN WARNA)
# ─────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Archivo+Black&family=Bungee&family=IBM+Plex+Mono:wght@400;600;700&family=Space+Grotesk:wght@400;500;600;700&display=swap');

/* ===== GLOBAL VARIABLES ===== */
:root {
    --blue: #2146FF;
    --yellow: #FFE500;
    --black: #000000;
    --white: #FFFFFF;
    --gray: #F2F2F2;
    --cyan: #00F5FF;
    --red: #FF2D2D;
    --lime: #B6FF00;
    --purple: #9B00FF;
    --border: 3px solid #000;
    --border-thick: 5px solid #000;
    --shadow: 6px 6px 0px #000;
    --shadow-sm: 4px 4px 0px #000;
    --shadow-lg: 8px 8px 0px #000;
    --shadow-hover: 2px 2px 0px #000;
}

html, body {
    font-family: 'Space Grotesk', sans-serif;
    background: #F2F2F2;
}

/* dot grid background */
.main {
    background-image: radial-gradient(circle, #00000018 1px, transparent 1px);
    background-size: 22px 22px;
    background-color: #F2F2F2;
}

.main .block-container {
    padding: 1.2rem 2rem 3rem;
    max-width: 1440px;
}

/* ===== STICKY HEADER ===== */
.sticky-header {
    position: fixed;
    top: 3.8rem;
    left: 18rem;
    right: 1rem;
    z-index: 999999;
    background: var(--yellow);
    border: var(--border-thick);
    box-shadow: var(--shadow);
    padding: 0.85rem 1.4rem;
    display: flex;
    align-items: center;
    gap: 12px;
}
.sticky-header-spacer { height: 100px; }
.sticky-header-icon {
    width: 40px;
    height: 40px;
    min-width: 40px;
    background: var(--blue);
    border: var(--border);
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 1.2rem;
    color: white;
}
.sticky-header-title {
    font-family: 'Archivo Black', sans-serif;
    font-size: 1.3rem;
    font-weight: 900;
    color: var(--black);
    text-transform: uppercase;
    letter-spacing: 0.02em;
    white-space: nowrap;
}
@media (max-width: 768px) {
    .sticky-header { left: 1rem; right: 1rem; top: 4rem; }
}

/* ===== SIDEBAR ===== */
[data-testid="stSidebar"] > div:first-child {
    background: var(--blue) !important;
    border-right: 5px solid var(--black) !important;
}
[data-testid="stSidebar"] * {
    color: var(--white) !important;
}
[data-testid="stSidebar"] label {
    color: var(--yellow) !important;
    font-family: 'Archivo Black', sans-serif !important;
    font-weight: 900 !important;
    letter-spacing: 0.1em !important;
    text-transform: uppercase !important;
    font-size: 0.7rem !important;
}
[data-testid="stSidebar"] .stRadio > div {
    gap: 4px;
    flex-direction: column;
}
[data-testid="stSidebar"] .stRadio label {
    background: rgba(0,0,0,0.25) !important;
    border: 2px solid rgba(255,255,255,0.3) !important;
    padding: 10px 14px !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-size: 0.85rem !important;
    font-weight: 600 !important;
    color: var(--white) !important;
    text-transform: none !important;
    letter-spacing: 0 !important;
    cursor: pointer;
    transition: all 0.15s;
    border-radius: 0 !important;
}
[data-testid="stSidebar"] .stRadio label:hover {
    background: var(--yellow) !important;
    border-color: var(--black) !important;
    color: var(--black) !important;
    box-shadow: 3px 3px 0 rgba(0,0,0,0.5) !important;
    transform: translate(-2px, -2px);
}
[data-testid="stSidebar"] .stRadio label[data-selected="true"],
[data-testid="stSidebar"] .stRadio label[aria-checked="true"] {
    background: var(--yellow) !important;
    border-color: var(--black) !important;
    color: var(--black) !important;
    font-weight: 700 !important;
}
[data-testid="stSidebar"] hr {
    border-color: rgba(255,255,255,0.3) !important;
    border-width: 2px !important;
}
[data-testid="stSidebar"] .stSlider * {
    color: var(--white) !important;
}
[data-testid="stSidebar"] .stButton > button {
    background: var(--yellow) !important;
    color: var(--black) !important;
    border: 3px solid var(--black) !important;
    box-shadow: 4px 4px 0 var(--black) !important;
    font-family: 'Archivo Black', sans-serif !important;
    font-weight: 900 !important;
    border-radius: 0 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.05em !important;
}
[data-testid="stSidebar"] .stButton > button:hover {
    transform: translate(3px, 3px) !important;
    box-shadow: 1px 1px 0 var(--black) !important;
    background: var(--white) !important;
}

/* ===== HERO SECTION ===== */
.hero-home {
    background: var(--blue);
    border: var(--border-thick);
    box-shadow: var(--shadow-lg);
    padding: 3rem 3rem 2.5rem;
    position: relative;
    overflow: hidden;
    margin-bottom: 2rem;
}
.hero-home::before {
    content: '';
    position: absolute;
    top: -40px; right: -40px;
    width: 200px; height: 200px;
    background: var(--yellow);
    border: 4px solid var(--black);
    transform: rotate(15deg);
    pointer-events: none;
    opacity: 0.35;
}
.hero-home::after {
    content: '';
    position: absolute;
    bottom: -30px; right: 120px;
    width: 120px; height: 120px;
    background: var(--cyan);
    border: 4px solid var(--black);
    border-radius: 50%;
    pointer-events: none;
    opacity: 0.3;
}
.hero-badge {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    background: var(--yellow);
    color: var(--black);
    border: 3px solid var(--black);
    box-shadow: 3px 3px 0 var(--black);
    padding: 5px 16px;
    font-family: 'IBM Plex Mono', monospace;
    font-size: 0.72rem;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.1em;
    margin-bottom: 1.4rem;
}
.hero-title {
    font-family: 'Archivo Black', sans-serif !important;
    font-size: 3rem !important;
    font-weight: 900 !important;
    color: var(--white) !important;
    line-height: 1.05 !important;
    margin: 0 0 1rem !important;
    text-transform: uppercase !important;
    letter-spacing: -0.5px !important;
}
.hero-sub {
    font-size: 1rem;
    color: rgba(255,255,255,0.88);
    max-width: 680px;
    line-height: 1.6;
    margin-bottom: 2rem;
    font-family: 'Space Grotesk', sans-serif;
}

/* ===== CARDS (BRUTALIST) ===== */
.modern-card {
    background: var(--white);
    border: var(--border-thick);
    box-shadow: var(--shadow);
    padding: 1.6rem;
    transition: transform 0.12s, box-shadow 0.12s;
    height: 100%;
    position: relative;
}
.modern-card::before {
    content: '';
    position: absolute;
    top: 6px; left: 6px;
    right: -6px; bottom: -6px;
    background: var(--yellow);
    z-index: -1;
    border: 2px solid var(--black);
}
.modern-card:hover {
    transform: translate(-3px, -3px);
    box-shadow: var(--shadow-lg);
}
.metric-icon {
    font-size: 1.9rem;
    margin-bottom: 0.6rem;
}
.metric-label {
    font-family: 'IBM Plex Mono', monospace;
    font-size: 0.68rem;
    font-weight: 700;
    color: #555;
    text-transform: uppercase;
    letter-spacing: 0.1em;
    margin-bottom: 4px;
}
.metric-value {
    font-family: 'Archivo Black', sans-serif;
    font-size: 2.4rem;
    font-weight: 900;
    color: var(--blue);
    line-height: 1;
}
.metric-desc {
    font-size: 0.78rem;
    color: #555;
    margin-top: 5px;
    font-family: 'IBM Plex Mono', monospace;
}

/* ===== PREDICTION CARD ===== */
.prediction-card {
    background: var(--white);
    border: var(--border-thick);
    box-shadow: var(--shadow);
    padding: 2rem;
    text-align: center;
}
.prediction-card.success {
    border-color: var(--black);
    background: #D1FAE5;
    color: #065F46;
}
.prediction-card.danger {
    border-color: var(--black);
    background: #FEE2E2;
    color: #991B1B;
}
.prediction-icon { 
    font-size: 3.5rem; 
    margin-bottom: 1rem; 
}
.prediction-result {
    font-family: 'Archivo Black', sans-serif;
    font-size: 2rem;
    font-weight: 900;
    text-transform: uppercase;
    margin-bottom: 0.5rem;
}
.prediction-prob { 
    font-size: 1rem; 
    color: #333; 
    margin-top: 1rem; 
}
.prediction-prob .prob-value {
    color: var(--black);
}
.badge-penerima {
    background: #D1FAE5;
    color: #065F46;
    border: 3px solid #065F46;
    box-shadow: 3px 3px 0 #065F46;
    padding: 5px 18px;
    font-family: 'Archivo Black', sans-serif;
    font-weight: 900;
    font-size: 0.85rem;
    text-transform: uppercase;
    display: inline-block;
}
.badge-tidak {
    background: #FEE2E2;
    color: #991B1B;
    border: 3px solid #991B1B;
    box-shadow: 3px 3px 0 #991B1B;
    padding: 5px 18px;
    font-family: 'Archivo Black', sans-serif;
    font-weight: 900;
    font-size: 0.85rem;
    text-transform: uppercase;
    display: inline-block;
}

/* ===== FORM & INPUT — NEO BRUTALISM WARM PALETTE ===== */

/* ── Base state: warm cream background ── */
.stTextInput input,
.stNumberInput input,
.stSelectbox > div > div,
textarea {
    background: #FFF8D6 !important;
    color: #111111 !important;
    border: 3px solid #000 !important;
    border-radius: 0 !important;
    box-shadow: 3px 3px 0 #000 !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-weight: 700 !important;
    transition: all 0.12s ease !important;
}

/* ── Hover state: slightly brighter cream ── */
.stTextInput input:hover,
.stNumberInput input:hover,
.stSelectbox > div > div:hover,
textarea:hover {
    background: #FFFBE6 !important;
    box-shadow: 4px 4px 0 #000 !important;
}

/* ── Focus state: bold yellow tactile feedback ── */
.stTextInput input:focus,
.stNumberInput input:focus,
.stSelectbox > div > div:focus-within,
textarea:focus {
    background: #FFE500 !important;
    border-color: #2146FF !important;
    box-shadow: 6px 6px 0 #000 !important;
    transform: translate(-2px, -2px) !important;
    outline: none !important;
}

/* ── Selectbox inner elements ── */
.stSelectbox div[data-baseweb="select"] {
    background: transparent !important;
}
.stSelectbox span {
    color: #ffffff !important;
    font-weight: 700 !important;
}

/* ── Placeholder: soft & lightweight ── */
.stTextInput input::placeholder,
.stNumberInput input::placeholder,
textarea::placeholder {
    color: #666666 !important;
    font-weight: 500 !important;
}

/* ── Number input stepper buttons ── */
.stNumberInput [data-testid="stNumberInputStepUp"],
.stNumberInput [data-testid="stNumberInputStepDown"] {
    background: #FFF8D6 !important;
    border: 2px solid #000 !important;
    border-radius: 0 !important;
    color: #111111 !important;
    transition: all 0.1s ease !important;
}
.stNumberInput [data-testid="stNumberInputStepUp"]:hover,
.stNumberInput [data-testid="stNumberInputStepDown"]:hover {
    background: #FFE500 !important;
}

/* ── File uploader ── */
.stFileUploader > div {
    border: 3px dashed var(--black) !important;
    border-radius: 0 !important;
    background: #F8F5E9 !important;
    box-shadow: 4px 4px 0 var(--black) !important;
    transition: all 0.12s ease !important;
}
.stFileUploader > div:hover {
    background: #FFF8D6 !important;
    box-shadow: 6px 6px 0 var(--black) !important;
}

/* ── Input label: uppercase mono style ── */
.stTextInput label,
.stNumberInput label,
.stSelectbox label {
    font-family: 'IBM Plex Mono', monospace !important;
    font-size: 0.7rem !important;
    font-weight: 700 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.08em !important;
    color: #333333 !important;
}

/* ===== ALERTS ===== */
.alert-success {
    background: #D1FAE5;
    border: 3px solid #065F46;
    box-shadow: 4px 4px 0 #065F46;
    padding: 12px 18px;
    color: #065F46;
    font-size: 0.9rem;
    font-weight: 600;
}
.alert-info {
    background: #DBEAFE;
    border: 3px solid #1E40AF;
    box-shadow: 4px 4px 0 #1E40AF;
    padding: 12px 18px;
    color: #1E40AF;
    font-size: 0.9rem;
    font-weight: 600;
}
.alert-warning {
    background: #FEF3C7;
    border: 3px solid #92400E;
    box-shadow: 4px 4px 0 #92400E;
    padding: 12px 18px;
    color: #92400E;
    font-size: 0.9rem;
    font-weight: 600;
}
.alert-error {
    background: #FEE2E2;
    border: 3px solid #991B1B;
    box-shadow: 4px 4px 0 #991B1B;
    padding: 12px 18px;
    color: #991B1B;
    font-size: 0.9rem;
    font-weight: 600;
}

/* ===== BUTTONS (GLOBAL) ===== */
.stButton > button {
    background: var(--blue) !important;
    color: var(--white) !important;
    border: 3px solid var(--black) !important;
    border-radius: 0 !important;
    box-shadow: 5px 5px 0 var(--black) !important;
    font-family: 'Archivo Black', sans-serif !important;
    font-weight: 900 !important;
    font-size: 0.85rem !important;
    letter-spacing: 0.05em !important;
    text-transform: uppercase !important;
    padding: 11px 22px !important;
    transition: transform 0.1s, box-shadow 0.1s !important;
}
.stButton > button:hover {
    transform: translate(4px, 4px) !important;
    box-shadow: 1px 1px 0 var(--black) !important;
    background: var(--yellow) !important;
    color: var(--black) !important;
}
.stButton > button:active {
    transform: translate(5px, 5px) !important;
    box-shadow: 0px 0px 0 var(--black) !important;
}
.stDownloadButton > button {
    background: var(--yellow) !important;
    color: var(--black) !important;
    border: 3px solid var(--black) !important;
    border-radius: 0 !important;
    box-shadow: 5px 5px 0 var(--black) !important;
    font-family: 'Archivo Black', sans-serif !important;
    font-weight: 900 !important;
    font-size: 0.82rem !important;
    letter-spacing: 0.05em !important;
    text-transform: uppercase !important;
    transition: transform 0.1s, box-shadow 0.1s !important;
}
.stDownloadButton > button:hover {
    transform: translate(4px, 4px) !important;
    box-shadow: 1px 1px 0 var(--black) !important;
}

/* ===== TABS ===== */
.stTabs [data-baseweb="tab-list"] {
    gap: 4px;
    background: var(--black);
    padding: 4px;
    border: 3px solid var(--black);
    box-shadow: 5px 5px 0 var(--black);
    border-radius: 0 !important;
}
.stTabs [data-baseweb="tab"] {
    border-radius: 0 !important;
    padding: 8px 16px !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-weight: 700 !important;
    font-size: 0.82rem !important;
    color: rgba(255,255,255,0.7) !important;
    text-transform: uppercase !important;
    letter-spacing: 0.05em !important;
    border: 2px solid transparent !important;
}
.stTabs [aria-selected="true"] {
    background: var(--yellow) !important;
    color: var(--black) !important;
    border: 2px solid var(--black) !important;
    font-weight: 900 !important;
}
.stTabs [data-baseweb="tab"]:hover {
    background: rgba(255,255,255,0.15) !important;
    color: var(--white) !important;
}

/* ===== PROGRESS ===== */
.stProgress > div > div > div {
    background: var(--blue) !important;
    border-radius: 0 !important;
    border: 1px solid var(--black) !important;
}
.stProgress > div > div {
    background: var(--gray) !important;
    border: 2px solid var(--black) !important;
    border-radius: 0 !important;
    box-shadow: 3px 3px 0 var(--black) !important;
}

/* ===== EXPANDER ===== */
.streamlit-expanderHeader {
    background: var(--blue) !important;
    border: 3px solid var(--black) !important;
    border-radius: 0 !important;
    box-shadow: 4px 4px 0 var(--black) !important;
    color: var(--white) !important;
    font-family: 'Archivo Black', sans-serif !important;
    font-weight: 900 !important;
    text-transform: uppercase !important;
}
.streamlit-expanderContent {
    border: 3px solid var(--black) !important;
    border-top: none !important;
    border-radius: 0 !important;
    background: var(--white) !important;
}

/* ===== DATAFRAME ===== */
.stDataFrame {
    border: 3px solid var(--black) !important;
    box-shadow: 5px 5px 0 var(--black) !important;
}
[data-testid="stDataFrame"] th {
    background: var(--blue) !important;
    color: var(--white) !important;
    font-family: 'Archivo Black', sans-serif !important;
    font-weight: 900 !important;
    text-transform: uppercase !important;
    border: 2px solid var(--black) !important;
}
[data-testid="stDataFrame"] td {
    border: 1px solid #ddd !important;
    font-family: 'IBM Plex Mono', monospace !important;
    font-size: 0.82rem !important;
    color: var(--black) !important;
}
[data-testid="stDataFrame"] tr:nth-child(even) td {
    background: rgba(33, 70, 255, 0.05) !important;
}
[data-testid="stDataFrame"] tr:hover td {
    background: rgba(255, 229, 0, 0.3) !important;
}

/* ===== METRIC WIDGET ===== */
[data-testid="metric-container"] {
    background: var(--white);
    border: 3px solid var(--black);
    box-shadow: var(--shadow);
    padding: 16px !important;
}
[data-testid="metric-container"] label {
    font-family: 'IBM Plex Mono', monospace !important;
    font-size: 0.7rem !important;
    text-transform: uppercase !important;
    letter-spacing: 0.08em !important;
    color: #555 !important;
}
[data-testid="metric-container"] [data-testid="stMetricValue"] {
    font-family: 'Archivo Black', sans-serif !important;
    color: var(--blue) !important;
    font-weight: 900 !important;
}

/* ===== FOOTER ===== */
.footer {
    background: var(--black);
    color: var(--white);
    border: var(--border-thick);
    box-shadow: var(--shadow-lg);
    padding: 2rem;
    text-align: center;
    margin-top: 3rem;
    position: relative;
    overflow: hidden;
}
.footer::before {
    content: '//BSM//';
    position: absolute;
    top: 50%; left: -10px;
    transform: translateY(-50%);
    font-family: 'Bungee', sans-serif;
    font-size: 5rem;
    color: rgba(255,229,0,0.08);
    white-space: nowrap;
    pointer-events: none;
    letter-spacing: 0.2em;
}
.footer-brand {
    font-family: 'Archivo Black', sans-serif;
    color: var(--yellow);
    font-weight: 900;
    font-size: 1.2rem;
    text-transform: uppercase;
    letter-spacing: 0.1em;
}
.footer-divider {
    width: 60px;
    height: 4px;
    background: var(--yellow);
    margin: 12px auto;
}

/* ===== INFO/WARNING/ERROR NATIVE ===== */
.stAlert {
    border: 3px solid var(--black) !important;
    border-radius: 0 !important;
    box-shadow: 4px 4px 0 var(--black) !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-weight: 600 !important;
}

/* ===== SCROLLBAR ===== */
::-webkit-scrollbar { width: 8px; height: 8px; }
::-webkit-scrollbar-track { background: var(--gray); border: 2px solid var(--black); }
::-webkit-scrollbar-thumb { background: var(--blue); border: 2px solid var(--black); }
::-webkit-scrollbar-thumb:hover { background: var(--yellow); border: 2px solid var(--black); }

/* ===== HEADINGS IN MAIN ===== */
.main h1, .main h2, .main h3 {
    font-family: 'Archivo Black', sans-serif !important;
    text-transform: uppercase !important;
    letter-spacing: 0.02em !important;
}
.main h3 {
    border-left: 5px solid var(--blue);
    padding-left: 10px;
    margin-top: 1.5rem;
}

/* ===== SPINNER ===== */
.stSpinner > div {
    border-top-color: var(--blue) !important;
}

/* ===== HIDE STREAMLIT BRANDING ===== */
#MainMenu { visibility: hidden; }
footer { visibility: hidden; }
.stDeployButton { display: none; }

/* ===== ANIMATIONS ===== */
@keyframes floatY {
    0%, 100% { transform: translateY(0px); }
    50% { transform: translateY(-8px); }
}
@keyframes fadeInUp {
    from { opacity: 0; transform: translateY(20px); }
    to   { opacity: 1; transform: translateY(0); }
}
.modern-card { animation: fadeInUp 0.3s ease both; }
.hero-home { animation: fadeInUp 0.4s ease both; }
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────────────────────
# SESSION STATE INIT
# ─────────────────────────────────────────────────────────────
if 'app_state' not in st.session_state:
    st.session_state.app_state = 'home'
if 'df_raw' not in st.session_state:
    st.session_state.df_raw = None
if 'df_processed' not in st.session_state:
    st.session_state.df_processed = None
if 'preprocess_info' not in st.session_state:
    st.session_state.preprocess_info = None
if 'train_results' not in st.session_state:
    st.session_state.train_results = None
if 'test_size' not in st.session_state:
    st.session_state.test_size = 0.2
if 'training_locked' not in st.session_state:
    st.session_state.training_locked = False
if 'batch_prediction_df' not in st.session_state:
    st.session_state.batch_prediction_df = None

# ─────────────────────────────────────────────────────────────
# HELPER FUNCTIONS (UI)
# ─────────────────────────────────────────────────────────────

def render_sticky_header(icon: str, title: str):
    """Render fixed header yang selalu terlihat di atas saat scroll."""
    st.markdown(f"""
    <div class="sticky-header">
        <div class="sticky-header-icon">{icon}</div>
        <div class="sticky-header-title">{title}</div>
    </div>
    <div class="sticky-header-spacer"></div>
    """, unsafe_allow_html=True)

def render_metric_card(icon: str, label: str, value: str, desc: str = ""):
    """Render single modern metric card."""
    st.markdown(f"""
    <div class="modern-card">
        <div class="metric-icon">{icon}</div>
        <div class="metric-label">{label}</div>
        <div class="metric-value">{value}</div>
        <div class="metric-desc">{desc}</div>
    </div>
    """, unsafe_allow_html=True)

def render_metric_cards(metrics: list, cols_per_row: int = 4):
    """Render multiple metric cards in a row."""
    cols = st.columns(cols_per_row)
    for i, m in enumerate(metrics):
        with cols[i % cols_per_row]:
            render_metric_card(m['icon'], m['label'], m['value'], m.get('desc', ''))

def alert(msg: str, kind: str = "success"):
    """Render alert box."""
    icons = {"success": "✅", "info": "ℹ️", "warning": "⚠️", "error": "❌"}
    st.markdown(f'<div class="alert-{kind}">{icons.get(kind, "•")} {msg}</div>', unsafe_allow_html=True)

def make_plotly_theme():
    """Return consistent Plotly theme dict — DARK MODE BLACK BACKGROUND."""
    return dict(
        paper_bgcolor="#111111",
        plot_bgcolor="#111111",
        font=dict(family="Space Grotesk, sans-serif", color="#E5E7EB", size=12),
        margin=dict(l=20, r=20, t=50, b=20),
        title_font=dict(family="Archivo Black, sans-serif", size=15, color="#FFFFFF"),
        legend_font=dict(color="#E5E7EB"),
        xaxis=dict(
            showgrid=True, 
            gridcolor="#333333", 
            linecolor="#555555", 
            linewidth=2,
            tickfont=dict(family="IBM Plex Mono, monospace", size=11, color="#BBBBBB"),
            title_font=dict(color="#E5E7EB")
        ),
        yaxis=dict(
            showgrid=True, 
            gridcolor="#333333", 
            linecolor="#555555", 
            linewidth=2,
            tickfont=dict(family="IBM Plex Mono, monospace", size=11, color="#BBBBBB"),
            title_font=dict(color="#E5E7EB")
        ),
        colorway=["#60A5FA", "#F87171", "#34D399", "#FBBF24", "#A78BFA", "#22D3EE"]
    )

def make_plotly_colors():
    """Return color scheme untuk visualisasi konsisten di dark background."""
    return {
        'penerima': '#34D399',           # Emerald green
        'tidak_penerima': '#F87171',     # Red
        'primary': '#60A5FA',             # Blue
        'secondary': '#A78BFA',           # Purple
        'accent': '#FBBF24',              # Amber
        'info': '#22D3EE',                # Cyan
        'gradient_blue': ['#1E3A5F', '#1E40AF', '#2563EB', '#3B82F6', '#60A5FA', '#93C5FD', '#BFDBFE'],
        'gradient_red': ['#5F1A1A', '#991B1B', '#DC2626', '#EF4444', '#F87171', '#FCA5A5', '#FECACA'],
        'gradient_green': ['#1A3F32', '#065F46', '#059669', '#10B981', '#34D399', '#6EE7B7', '#A7F3D0'],
        'pie_colors': ['#60A5FA', '#F87171', '#34D399', '#FBBF24', '#A78BFA'],
        'bar_blue': '#3B82F6',
        'bar_red': '#EF4444',
        'bar_green': '#10B981',
    }

# ─────────────────────────────────────────────────────────────
# PREPROCESSING ENGINE
# ─────────────────────────────────────────────────────────────

def print_preprocessing_info(df, info):
    """Cetak informasi preprocessing ke terminal."""
    print("\n" + "="*80)
    print("PREPROCESSING DATA")
    print("="*80)
    
    print(f"\n[1] DATA AWAL")
    print(f"    Jumlah data awal : {info['row_after_dup'] + info['dup_removed']} baris")
    
    print(f"\n[2] MISSING VALUE")
    print(f"    Missing value sebelum penanganan:")
    for col, missing in info['missing_before'].items():
        if missing > 0:
            print(f"    - {col}: {missing}")
    
    print(f"\n    Missing value setelah penanganan:")
    for col, missing in info['missing_after'].items():
        if missing > 0:
            print(f"    - {col}: {missing}")
    
    print(f"\n[3] DUPLIKASI DATA")
    print(f"    Data duplikat dihapus : {info['dup_removed']} baris")
    print(f"    Data setelah dedup   : {info['row_after_dup']} baris")
    
    print(f"\n[4] MAPPING ENCODING")
    
    le = info['le_pekerjaan']
    print(f"\n    Pekerjaan Orang Tua → Label Encoding:")
    for kd, pk in zip(le.transform(le.classes_), le.classes_):
        print(f"    {pk:30s} → {kd}")
    
    print(f"\n    Status Rumah → Binary Encoding:")
    print(f"    {'Kontrak/sewa':30s} → 0")
    print(f"    {'Milik Sendiri':30s} → 1")
    
    print(f"\n    Label → Binary Encoding:")
    print(f"    {'Tidak':30s} → 0")
    print(f"    {'Ya':30s} → 1")
    
    print(f"\n[5] STANDARDISASI Z-SCORE (PENDAPATAN ORANG TUA)")
    scaler = info['scaler']
    print(f"    Mean   : {scaler.mean_[0]:.6f}")
    print(f"    Std    : {scaler.scale_[0]:.6f}")
    
    print(f"\n[6] 5 BARIS PERTAMA HASIL PREPROCESSING")
    cols_show = ['NO', 'NAMA PESERTA DIDIK', 'KELAS', 'KELAS_NUM',
                 'PENDAPATAN ORANG TUA', 'PENDAPATAN_ZSCORE',
                 'PEKERJAAN ORANG TUA', 'PEKERJAAN ORANG TUA_ENC',
                 'JUMLAH TANGGUNGAN', 'STATUS RUMAH', 'STATUS RUMAH_ENC',
                 'LABEL', 'LABEL_ENC']
    print(df[cols_show].head(5).to_string(index=False))
    
    print("\n" + "="*80)
    print("PREPROCESSING SELESAI")
    print("="*80 + "\n")

def run_preprocessing(df_raw: pd.DataFrame):
    """
    Jalankan seluruh pipeline preprocessing.
    Kembalikan (df_processed, info_dict).
    """
    df = df_raw.copy()

    expected_cols = [
        'NO', 'NAMA PESERTA DIDIK', 'SEKOLAH', 'KELAS',
        'NIK', 'NISN', 'NAMA AYAH/IBU', 'KECAMATAN',
        'BESARAN BIAYA', 'PENDAPATAN ORANG TUA',
        'PEKERJAAN ORANG TUA', 'JUMLAH TANGGUNGAN',
        'STATUS RUMAH', 'LABEL'
    ]
    if len(df.columns) == len(expected_cols):
        df.columns = expected_cols

    info = {}
    info['missing_before'] = df.isnull().sum().to_dict()

    df['PENDAPATAN ORANG TUA'] = pd.to_numeric(df['PENDAPATAN ORANG TUA'], errors='coerce')
    df['JUMLAH TANGGUNGAN']    = pd.to_numeric(df['JUMLAH TANGGUNGAN'], errors='coerce')
    df['PENDAPATAN ORANG TUA'] = df['PENDAPATAN ORANG TUA'].fillna(df['PENDAPATAN ORANG TUA'].mean())
    df['JUMLAH TANGGUNGAN']    = df['JUMLAH TANGGUNGAN'].fillna(df['JUMLAH TANGGUNGAN'].median())

    info['missing_after'] = df.isnull().sum().to_dict()
    before_dup = len(df)
    df = df.drop_duplicates()
    after_dup = len(df)
    info['dup_removed'] = before_dup - after_dup
    info['row_after_dup'] = after_dup

    df['KELAS_NUM'] = df['KELAS'].astype(str).str.extract(r'(\d+)')
    df['KELAS_NUM'] = pd.to_numeric(df['KELAS_NUM'], errors='coerce').fillna(0).astype(int)

    le_pekerjaan = LabelEncoder()
    df['PEKERJAAN ORANG TUA_ENC'] = le_pekerjaan.fit_transform(
        df['PEKERJAAN ORANG TUA'].astype(str)
    )

    status_rumah_mapping = {"Milik Sendiri": 1, "Kontrak/sewa": 0}
    df['STATUS RUMAH_ENC'] = df['STATUS RUMAH'].map(status_rumah_mapping)
    df['STATUS RUMAH_ENC'] = df['STATUS RUMAH_ENC'].fillna(0).astype(int)

    scaler = StandardScaler()
    df['PENDAPATAN_ZSCORE'] = scaler.fit_transform(df[['PENDAPATAN ORANG TUA']])

    df['LABEL'] = df['LABEL'].astype(str).str.strip().str.capitalize()
    label_mapping = {"Ya": 1, "Tidak": 0}
    df['LABEL_ENC'] = df['LABEL'].map(label_mapping)

    info['le_pekerjaan'] = le_pekerjaan
    info['scaler']       = scaler
    info['total_rows']   = len(df)

    print_preprocessing_info(df, info)
    return df, info

# ─────────────────────────────────────────────────────────────
# TRAINING ENGINE
# ─────────────────────────────────────────────────────────────

def run_training(df: pd.DataFrame, test_size: float = 0.2):
    """
    Latih model Gaussian Naive Bayes.
    Kembalikan dictionary berisi semua artefak.
    """
    feature_cols = [
        'KELAS_NUM',
        'PENDAPATAN_ZSCORE',
        'PEKERJAAN ORANG TUA_ENC',
        'JUMLAH TANGGUNGAN',
        'STATUS RUMAH_ENC'
    ]

    X = df[feature_cols].copy()
    y = df['LABEL_ENC'].copy()

    valid_idx = X.notna().all(axis=1) & y.notna()
    X = X[valid_idx]
    y = y[valid_idx]

    print("\n" + "="*80)
    print("PEMBAGIAN DATA TRAINING DAN TESTING")
    print("="*80)
    
    if test_size == 0.0:
        X_train, X_test = X.copy(), pd.DataFrame(columns=X.columns)
        y_train, y_test = y.copy(), pd.Series(dtype=y.dtype)
        
        print(f"\n    Total data valid : {len(X)}")
        print(f"    Data Training    : {len(X_train)} (100.0%)")
        print(f"    Data Testing     : 0 (0.0%)")
        print(f"\n    ⚠️ Mode Training Only - Semua data untuk training")
        print(f"    ⚠️ Tidak ada data testing otomatis")
        print("\n" + "="*80 + "\n")
    elif test_size == 1.0:
        print(f"\n    Total data valid : {len(X)}")
        print(f"    ⚠️ Testing 100% - Semua data testing, tidak bisa training")
        print(f"    ⚠️ Harap upload data training manual")
        print("\n" + "="*80 + "\n")
        
        return {
            'model': None,
            'X': X, 'y': y,
            'X_train': pd.DataFrame(columns=X.columns),
            'X_test': X,
            'y_train': pd.Series(dtype=y.dtype),
            'y_test': y,
            'y_pred': np.array([]),
            'y_pred_proba': np.array([]),
            'train_acc': None,
            'test_acc': None,
            'cm': np.array([[0, 0], [0, 0]]),
            'prec0': 0, 'rec0': 0, 'f1_0': 0,
            'prec1': 0, 'rec1': 0, 'f1_1': 0,
            'macro_prec': 0, 'macro_rec': 0, 'macro_f1': 0,
            'w_prec': 0, 'w_rec': 0, 'w_f1': 0,
            'sup0': 0, 'sup1': 0,
            'feature_cols': feature_cols,
            'results_df': pd.DataFrame(columns=[
                'No', 'Actual', 'Predicted', 'Prob Tidak (%)', 'Prob Ya (%)', 'Status'
            ])
        }
    else:
        X_train, X_test, y_train, y_test = train_test_split(
            X, y, test_size=test_size, random_state=42, stratify=y
        )
        
        print(f"\n    Total data valid : {len(X)}")
        print(f"    Data Training    : {len(X_train)} ({(len(X_train)/len(X))*100:.1f}%)")
        print(f"    Data Testing     : {len(X_test)} ({(len(X_test)/len(X))*100:.1f}%)")
        
        train_dist = y_train.value_counts()
        test_dist = y_test.value_counts()
        
        print(f"\n    Distribusi Kelas pada Data Training:")
        print(f"    - Tidak Penerima (0) : {train_dist.get(0, 0)}")
        print(f"    - Penerima (1)       : {train_dist.get(1, 0)}")
        
        print(f"\n    Distribusi Kelas pada Data Testing:")
        print(f"    - Tidak Penerima (0) : {test_dist.get(0, 0)}")
        print(f"    - Penerima (1)       : {test_dist.get(1, 0)}")
        print("\n" + "="*80 + "\n")

    if test_size < 1.0:
        model = GaussianNB()
        model.fit(X_train, y_train)

        print("\n" + "="*80)
        print("TRAINING GAUSSIAN NAIVE BAYES")
        print("="*80)
        
        print(f"\n[1] PRIOR PROBABILITY")
        print(f"    P(Tidak) = {model.class_prior_[0]:.6f} ({model.class_prior_[0]*100:.2f}%)")
        print(f"    P(Ya)    = {model.class_prior_[1]:.6f} ({model.class_prior_[1]*100:.2f}%)")
        
        print(f"\n[2] MEAN (THETA_) PER FITUR")
        print(f"    {'Fitur':35s} {'Tidak (0)':>12s} {'Ya (1)':>12s}")
        print(f"    {'-'*35} {'-'*12} {'-'*12}")
        for i, feat in enumerate(feature_cols):
            print(f"    {feat:35s} {model.theta_[0][i]:12.6f} {model.theta_[1][i]:12.6f}")
        
        print(f"\n[3] VARIANCE (VAR_) PER FITUR")
        print(f"    {'Fitur':35s} {'Tidak (0)':>12s} {'Ya (1)':>12s}")
        print(f"    {'-'*35} {'-'*12} {'-'*12}")
        for i, feat in enumerate(feature_cols):
            print(f"    {feat:35s} {model.var_[0][i]:12.6f} {model.var_[1][i]:12.6f}")
        
        print(f"\n[4] STANDARD DEVIATION (SIGMA) PER FITUR")
        print(f"    {'Fitur':35s} {'Tidak (0)':>12s} {'Ya (1)':>12s}")
        print(f"    {'-'*35} {'-'*12} {'-'*12}")
        for i, feat in enumerate(feature_cols):
            print(f"    {feat:35s} {np.sqrt(model.var_[0][i]):12.6f} {np.sqrt(model.var_[1][i]):12.6f}")
        
        print("\n" + "="*80)
        print("TRAINING SELESAI")
        print("="*80 + "\n")

        y_train_pred = model.predict(X_train)
        train_acc = accuracy_score(y_train, y_train_pred)
    else:
        model = None
        train_acc = None

    if test_size == 0.0:
        print("\n" + "="*80)
        print("EVALUASI MODEL")
        print("="*80)
        print(f"\n    Akurasi Training : {train_acc:.6f} ({train_acc*100:.2f}%)")
        print(f"\n    ⚠️ Testing 0% - Tidak ada data testing untuk evaluasi")
        print("\n" + "="*80 + "\n")
        
        results_dict = {
            'model': model,
            'X': X, 'y': y,
            'X_train': X_train, 'X_test': X_test,
            'y_train': y_train, 'y_test': y_test,
            'y_pred': np.array([]),
            'y_pred_proba': np.array([]),
            'train_acc': train_acc,
            'test_acc': None,
            'cm': np.array([[0, 0], [0, 0]]),
            'prec0': 0, 'rec0': 0, 'f1_0': 0,
            'prec1': 0, 'rec1': 0, 'f1_1': 0,
            'macro_prec': 0, 'macro_rec': 0, 'macro_f1': 0,
            'w_prec': 0, 'w_rec': 0, 'w_f1': 0,
            'sup0': 0, 'sup1': 0,
            'feature_cols': feature_cols,
            'results_df': pd.DataFrame(columns=[
                'No', 'Actual', 'Predicted', 'Prob Tidak (%)', 'Prob Ya (%)', 'Status'
            ])
        }
    elif test_size == 1.0:
        print("\n" + "="*80)
        print("EVALUASI MODEL")
        print("="*80)
        print(f"\n    ⚠️ Testing 100% - Tidak ada model yang dilatih")
        print(f"    ⚠️ Silakan upload data training manual di Dashboard")
        print("\n" + "="*80 + "\n")
        
        results_dict = {
            'model': None,
            'X': X, 'y': y,
            'X_train': X_train, 'X_test': X_test,
            'y_train': y_train, 'y_test': y_test,
            'y_pred': np.array([]),
            'y_pred_proba': np.array([]),
            'train_acc': None,
            'test_acc': None,
            'cm': np.array([[0, 0], [0, 0]]),
            'prec0': 0, 'rec0': 0, 'f1_0': 0,
            'prec1': 0, 'rec1': 0, 'f1_1': 0,
            'macro_prec': 0, 'macro_rec': 0, 'macro_f1': 0,
            'w_prec': 0, 'w_rec': 0, 'w_f1': 0,
            'sup0': 0, 'sup1': 0,
            'feature_cols': feature_cols,
            'results_df': pd.DataFrame(columns=[
                'No', 'Actual', 'Predicted', 'Prob Tidak (%)', 'Prob Ya (%)', 'Status'
            ])
        }
    else:
        y_pred = model.predict(X_test)
        y_pred_proba = model.predict_proba(X_test)
        test_acc = accuracy_score(y_test, y_pred)
        cm = confusion_matrix(y_test, y_pred)

        print("\n" + "="*80)
        print("PROSES PREDIKSI PER DATA TESTING")
        print("="*80)
        
        X_test_array = X_test.values
        y_test_array = y_test.values
        
        for idx in range(len(X_test_array)):
            print(f"\n{'─'*60}")
            print(f"DATA TEST KE-{idx+1}")
            print(f"{'─'*60}")
            
            x = X_test_array[idx]
            actual = y_test_array[idx]
            
            print(f"\n    Fitur:")
            for i, feat in enumerate(feature_cols):
                print(f"    - {feat} = {x[i]:.6f}")
            
            theta = model.theta_
            var = model.var_
            class_prior = model.class_prior_
            
            posteriors = []
            for c in range(2):
                log_prior = np.log(class_prior[c])
                log_likelihood = 0
                for i in range(len(feature_cols)):
                    coef = 1.0 / np.sqrt(2 * np.pi * var[c][i])
                    exponent = np.exp(-((x[i] - theta[c][i]) ** 2) / (2 * var[c][i]))
                    likelihood = coef * exponent
                    log_likelihood += np.log(max(likelihood, 1e-300))
                
                posterior_log = log_prior + log_likelihood
                posteriors.append(posterior_log)
            
            max_log = max(posteriors)
            exp_post = [np.exp(p - max_log) for p in posteriors]
            total = sum(exp_post)
            prob = [e / total for e in exp_post]
            
            print(f"\n    Posterior Log:")
            print(f"    - Log P(Tidak) = {posteriors[0]:.6f}")
            print(f"    - Log P(Ya)    = {posteriors[1]:.6f}")
            
            print(f"\n    Probabilitas Akhir:")
            print(f"    - Prob Tidak = {prob[0]*100:.2f}%")
            print(f"    - Prob Ya    = {prob[1]*100:.2f}%")
            
            pred = np.argmax(prob)
            pred_label = "Ya" if pred == 1 else "Tidak"
            actual_label = "Ya" if actual == 1 else "Tidak"
            status = "BENAR" if pred == actual else "SALAH"
            
            print(f"\n    Hasil:")
            print(f"    - Prediksi = {pred_label}")
            print(f"    - Aktual   = {actual_label}")
            print(f"    - Status   = {status}")
        
        print(f"\n{'─'*60}")
        print("\n" + "="*80)
        print("PROSES PREDIKSI SELESAI")
        print("="*80 + "\n")

        TP = cm[1, 1]; TN = cm[0, 0]
        FP = cm[0, 1]; FN = cm[1, 0]
        
        prec0 = TN / (TN + FN) if (TN + FN) > 0 else 0
        rec0  = TN / (TN + FP) if (TN + FP) > 0 else 0
        f1_0  = 2 * prec0 * rec0 / (prec0 + rec0) if (prec0 + rec0) > 0 else 0
        
        prec1 = TP / (TP + FP) if (TP + FP) > 0 else 0
        rec1  = TP / (TP + FN) if (TP + FN) > 0 else 0
        f1_1  = 2 * prec1 * rec1 / (prec1 + rec1) if (prec1 + rec1) > 0 else 0
        
        sup0 = int((y_test == 0).sum())
        sup1 = int((y_test == 1).sum())
        tot  = sup0 + sup1
        
        macro_prec = (prec0 + prec1) / 2
        macro_rec  = (rec0  + rec1)  / 2
        macro_f1   = (f1_0  + f1_1)  / 2
        w_prec     = (prec0 * sup0 + prec1 * sup1) / tot if tot > 0 else 0
        w_rec      = (rec0  * sup0 + rec1  * sup1) / tot if tot > 0 else 0
        w_f1       = (f1_0  * sup0 + f1_1  * sup1) / tot if tot > 0 else 0

        print("\n" + "="*80)
        print("EVALUASI MODEL")
        print("="*80)
        
        print(f"\n[1] CONFUSION MATRIX")
        print(f"    [[TN  FP]     [[{cm[0][0]:4d}  {cm[0][1]:4d}]")
        print(f"     [FN  TP]]  =  [{cm[1][0]:4d}  {cm[1][1]:4d}]]")
        
        print(f"\n    True Positive  (TP) : {TP}")
        print(f"    True Negative  (TN) : {TN}")
        print(f"    False Positive (FP) : {FP}")
        print(f"    False Negative (FN) : {FN}")
        
        print(f"\n[2] METRIK EVALUASI")
        print(f"\n    Accuracy : {test_acc:.6f} ({test_acc*100:.2f}%)")
        
        print(f"\n    Kelas Tidak Penerima (0):")
        print(f"    - Precision : {prec0:.6f} ({prec0*100:.2f}%)")
        print(f"    - Recall    : {rec0:.6f} ({rec0*100:.2f}%)")
        print(f"    - F1-Score  : {f1_0:.6f} ({f1_0*100:.2f}%)")
        
        print(f"\n    Kelas Penerima (1):")
        print(f"    - Precision : {prec1:.6f} ({prec1*100:.2f}%)")
        print(f"    - Recall    : {rec1:.6f} ({rec1*100:.2f}%)")
        print(f"    - F1-Score  : {f1_1:.6f} ({f1_1*100:.2f}%)")
        
        print(f"\n    Macro Average:")
        print(f"    - Precision : {macro_prec:.6f} ({macro_prec*100:.2f}%)")
        print(f"    - Recall    : {macro_rec:.6f} ({macro_rec*100:.2f}%)")
        print(f"    - F1-Score  : {macro_f1:.6f} ({macro_f1*100:.2f}%)")
        
        print(f"\n    Weighted Average:")
        print(f"    - Precision : {w_prec:.6f} ({w_prec*100:.2f}%)")
        print(f"    - Recall    : {w_rec:.6f} ({w_rec*100:.2f}%)")
        print(f"    - F1-Score  : {w_f1:.6f} ({w_f1*100:.2f}%)")
        
        benar = int(TP + TN)
        salah = int(FP + FN)
        print(f"\n[3] RINGKASAN AKHIR")
        print(f"    Total prediksi benar : {benar} dari {tot} data")
        print(f"    Total prediksi salah : {salah} dari {tot} data")
        print(f"    Akurasi akhir        : {test_acc*100:.2f}%")
        
        print("\n" + "="*80)
        print("EVALUASI SELESAI")
        print("="*80 + "\n")

        results_dict = {
            'model': model,
            'X': X, 'y': y,
            'X_train': X_train, 'X_test': X_test,
            'y_train': y_train, 'y_test': y_test,
            'y_pred': y_pred,
            'y_pred_proba': y_pred_proba,
            'train_acc': train_acc,
            'test_acc': test_acc,
            'cm': cm,
            'prec0': prec0, 'rec0': rec0,  'f1_0': f1_0,
            'prec1': prec1, 'rec1': rec1,  'f1_1': f1_1,
            'macro_prec': macro_prec, 'macro_rec': macro_rec, 'macro_f1': macro_f1,
            'w_prec': w_prec, 'w_rec': w_rec, 'w_f1': w_f1,
            'sup0': sup0, 'sup1': sup1,
            'feature_cols': feature_cols,
            'results_df': pd.DataFrame({
                'No': range(1, len(y_test) + 1),
                'Actual':    ['Penerima' if v == 1 else 'Tidak Penerima' for v in y_test.values],
                'Predicted': ['Penerima' if v == 1 else 'Tidak Penerima' for v in y_pred],
                'Prob Tidak (%)': (y_pred_proba[:, 0] * 100).round(2),
                'Prob Ya (%)':    (y_pred_proba[:, 1] * 100).round(2),
                'Status': ['✓ Benar' if a == p else '✗ Salah'
                           for a, p in zip(y_test.values, y_pred)]
            })
        }

    return results_dict

# ─────────────────────────────────────────────────────────────
# EXCEL EXPORT HELPERS
# ─────────────────────────────────────────────────────────────

def _safe_excel_value(val):
    """Konversi nilai pandas NA menjadi None untuk kompatibilitas Excel."""
    if pd.isna(val):
        return None
    return val

def _style_header(cell, bg="065f46", fg="FFFFFF"):
    cell.font      = Font(name='Inter', bold=True, color=fg, size=11)
    cell.fill      = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='BFBFBF')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

def _style_data(cell, bg="FFFFFF", bold=False, align='center'):
    cell.font      = Font(name='Inter', size=10, bold=bold)
    cell.fill      = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal=align, vertical='center')
    thin = Side(style='thin', color='D9D9D9')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

def export_preprocessing_excel(df: pd.DataFrame, le_pekerjaan: LabelEncoder) -> bytes:
    """Export data preprocessing ke Excel dengan styling."""
    df_exp = df[[
        'NO', 'NAMA PESERTA DIDIK', 'SEKOLAH', 'KELAS', 'KELAS_NUM',
        'PENDAPATAN ORANG TUA', 'PENDAPATAN_ZSCORE',
        'PEKERJAAN ORANG TUA', 'PEKERJAAN ORANG TUA_ENC',
        'JUMLAH TANGGUNGAN', 'STATUS RUMAH', 'STATUS RUMAH_ENC',
        'LABEL', 'LABEL_ENC'
    ]].copy()

    for col in ['KELAS_NUM', 'PEKERJAAN ORANG TUA_ENC', 'STATUS RUMAH_ENC', 'LABEL_ENC']:
        df_exp[col] = pd.to_numeric(df_exp[col], errors='coerce')

    df_exp.columns = [
        'No', 'Nama Peserta Didik', 'Sekolah', 'Kelas', 'Kelas (Angka)',
        'Pendapatan Orang Tua (Rp)', 'Pendapatan (Z-Score)',
        'Pekerjaan Orang Tua', 'Pekerjaan (Encoded)',
        'Jumlah Tanggungan', 'Status Rumah', 'Status Rumah (Encoded)',
        'Label', 'Label (Encoded)'
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Preprocessing"

    ws.merge_cells('A1:N1')
    tc = ws['A1']
    tc.value = "DATA HASIL PREPROCESSING - BSM 2022-2024"
    tc.font = Font(name='Inter', bold=True, size=14, color='FFFFFF')
    tc.fill = PatternFill("solid", start_color="065f46")
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:N2')
    sc = ws['A2']
    sc.value = f"Total Data: {len(df_exp)} baris | Fitur: 5 | Target: LABEL (Ya=1, Tidak=0)"
    sc.font  = Font(name='Inter', size=9, italic=True, color='595959')
    sc.fill  = PatternFill("solid", start_color="d1fae5")
    sc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18

    headers = list(df_exp.columns)
    for ci, h in enumerate(headers, 1):
        _style_header(ws.cell(row=3, column=ci, value=h))
    ws.row_dimensions[3].height = 32

    row_colors = ["FFFFFF", "ecfdf5"]
    for ri, row in enumerate(df_exp.itertuples(index=False), 4):
        bg = row_colors[(ri - 4) % 2]
        for ci, val in enumerate(row, 1):
            safe_val = _safe_excel_value(val)
            cell = ws.cell(row=ri, column=ci, value=safe_val)
            _style_data(cell, bg)
            if ci == 6:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif ci == 7:
                cell.number_format = '0.000000'
                cell.alignment = Alignment(horizontal='right', vertical='center')

    for i, w in enumerate([5, 28, 28, 10, 12, 22, 18, 22, 16, 16, 18, 18, 8, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wm = wb.create_sheet("Mapping Encoding")
    wm.merge_cells('A1:C1')
    wm['A1'].value = "MAPPING ENCODING"
    wm['A1'].font  = Font(name='Inter', bold=True, size=13, color='FFFFFF')
    wm['A1'].fill  = PatternFill("solid", start_color="065f46")
    wm['A1'].alignment = Alignment(horizontal='center', vertical='center')
    wm.row_dimensions[1].height = 26

    _style_header(wm.cell(row=3, column=1, value="Pekerjaan Orang Tua"), "059669")
    _style_header(wm.cell(row=3, column=2, value="Kode"), "059669")
    for i, (kd, pk) in enumerate(zip(le_pekerjaan.transform(le_pekerjaan.classes_), le_pekerjaan.classes_), 4):
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(wm.cell(row=i, column=1, value=pk), bg, align='left')
        _style_data(wm.cell(row=i, column=2, value=int(kd)), bg)

    rs = 4 + len(le_pekerjaan.classes_) + 2
    _style_header(wm.cell(row=rs, column=1, value="Status Rumah"), "059669")
    _style_header(wm.cell(row=rs, column=2, value="Kode"), "059669")
    for i, (s, k) in enumerate([("Kontrak/sewa", 0), ("Milik Sendiri", 1)], rs + 1):
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(wm.cell(row=i, column=1, value=s), bg, align='left')
        _style_data(wm.cell(row=i, column=2, value=k), bg)

    rs2 = rs + 4
    _style_header(wm.cell(row=rs2, column=1, value="Label"), "059669")
    _style_header(wm.cell(row=rs2, column=2, value="Kode"), "059669")
    for i, (lb, k) in enumerate([("Tidak", 0), ("Ya", 1)], rs2 + 1):
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(wm.cell(row=i, column=1, value=lb), bg, align='left')
        _style_data(wm.cell(row=i, column=2, value=k), bg)

    wm.column_dimensions['A'].width = 26
    wm.column_dimensions['B'].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def export_normalisasi_excel(df: pd.DataFrame, info: dict) -> bytes:
    """Export hasil normalisasi ke Excel."""
    scaler = info['scaler']
    mean_val = scaler.mean_[0]
    std_val = scaler.scale_[0]
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "Hasil Normalisasi"
    
    _style_header(ws1.cell(row=1, column=1, value="No"))
    _style_header(ws1.cell(row=1, column=2, value="Nama Peserta Didik"))
    _style_header(ws1.cell(row=1, column=3, value="Pendapatan Asli (Rp)"))
    _style_header(ws1.cell(row=1, column=4, value="Pendapatan Z-Score"))
    _style_header(ws1.cell(row=1, column=5, value="Rumus Z-Score"))
    
    for i, (_, row) in enumerate(df.iterrows(), 2):
        z = (row['PENDAPATAN ORANG TUA'] - mean_val) / std_val
        rumus = f"({row['PENDAPATAN ORANG TUA']} - {mean_val:.2f}) / {std_val:.2f}"
        
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(ws1.cell(row=i, column=1, value=_safe_excel_value(row['NO'])), bg)
        _style_data(ws1.cell(row=i, column=2, value=_safe_excel_value(row['NAMA PESERTA DIDIK'])), bg, align='left')
        _style_data(ws1.cell(row=i, column=3, value=_safe_excel_value(row['PENDAPATAN ORANG TUA'])), bg)
        ws1.cell(row=i, column=3).number_format = '#,##0'
        _style_data(ws1.cell(row=i, column=4, value=round(z, 6)), bg)
        _style_data(ws1.cell(row=i, column=5, value=rumus), bg, align='left')
    
    ws1.column_dimensions['A'].width = 6
    ws1.column_dimensions['B'].width = 30
    ws1.column_dimensions['C'].width = 22
    ws1.column_dimensions['D'].width = 18
    ws1.column_dimensions['E'].width = 35
    
    ws2 = wb.create_sheet("Parameter Normalisasi")
    _style_header(ws2.cell(row=1, column=1, value="Parameter"), "059669")
    _style_header(ws2.cell(row=1, column=2, value="Nilai"), "059669")
    _style_header(ws2.cell(row=1, column=3, value="Keterangan"), "059669")
    
    _style_data(ws2.cell(row=2, column=1, value="Mean (μ)"), "FFFFFF", align='left')
    _style_data(ws2.cell(row=2, column=2, value=round(mean_val, 4)), "FFFFFF")
    _style_data(ws2.cell(row=2, column=3, value="Rata-rata pendapatan orang tua"), "FFFFFF", align='left')
    
    _style_data(ws2.cell(row=3, column=1, value="Standard Deviation (σ)"), "ecfdf5", align='left')
    _style_data(ws2.cell(row=3, column=2, value=round(std_val, 4)), "ecfdf5")
    _style_data(ws2.cell(row=3, column=3, value="Simpangan baku pendapatan orang tua"), "ecfdf5", align='left')
    
    ws2.merge_cells('A5:C5')
    _style_header(ws2.cell(row=5, column=1, value="Rumus Z-Score: z = (x - μ) / σ"), "047857")
    
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 35
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ─────────────────────────────────────────────────────────────
# FITUR BARU: PREDIKSI DATA BARU
# ─────────────────────────────────────────────────────────────

def preprocess_single_input(kelas, pendapatan, pekerjaan, tanggungan, status_rumah):
    """
    Preprocess single input untuk prediksi manual.
    """
    le_pekerjaan = st.session_state.preprocess_info['le_pekerjaan']
    scaler = st.session_state.preprocess_info['scaler']
    
    pekerjaan_enc = le_pekerjaan.transform([pekerjaan])[0]
    
    status_rumah_mapping = {"Milik Sendiri": 1, "Kontrak/sewa": 0}
    status_rumah_enc = status_rumah_mapping.get(status_rumah, 0)
    
    pendapatan_zscore = (pendapatan - scaler.mean_[0]) / scaler.scale_[0]
    
    features = np.array([
        int(kelas),
        pendapatan_zscore,
        pekerjaan_enc,
        int(tanggungan),
        status_rumah_enc
    ]).reshape(1, -1)
    
    return features

def predict_single(model, features):
    """Melakukan prediksi untuk satu data."""
    prediction = model.predict(features)[0]
    probabilities = model.predict_proba(features)[0]
    return prediction, probabilities

def preprocess_batch_data(df_input):
    """Preprocess data batch untuk prediksi massal."""
    required_columns = ['KELAS', 'PENDAPATAN ORANG TUA', 'PEKERJAAN ORANG TUA', 
                       'JUMLAH TANGGUNGAN', 'STATUS RUMAH']
    
    missing_cols = [col for col in required_columns if col not in df_input.columns]
    if missing_cols:
        raise ValueError(f"Kolom berikut tidak ditemukan: {', '.join(missing_cols)}")
    
    le_pekerjaan = st.session_state.preprocess_info['le_pekerjaan']
    scaler = st.session_state.preprocess_info['scaler']
    
    df_processed = df_input.copy()
    
    df_processed['KELAS_NUM'] = df_processed['KELAS'].astype(str).str.extract(r'(\d+)')
    df_processed['KELAS_NUM'] = pd.to_numeric(df_processed['KELAS_NUM'], errors='coerce').fillna(0).astype(int)
    
    pekerjaan_valid = set(le_pekerjaan.classes_)
    pekerjaan_input = set(df_processed['PEKERJAAN ORANG TUA'].astype(str).unique())
    pekerjaan_unknown = pekerjaan_input - pekerjaan_valid
    
    if pekerjaan_unknown:
        raise ValueError(
            f"Kategori pekerjaan tidak dikenal: {', '.join(pekerjaan_unknown)}. "
            f"Kategori yang valid: {', '.join(sorted(pekerjaan_valid))}"
        )
    
    df_processed['PEKERJAAN_ENC'] = le_pekerjaan.transform(
        df_processed['PEKERJAAN ORANG TUA'].astype(str)
    )
    
    status_rumah_mapping = {"Milik Sendiri": 1, "Kontrak/sewa": 0}
    df_processed['STATUS_RUMAH_ENC'] = df_processed['STATUS RUMAH'].map(status_rumah_mapping)
    df_processed['STATUS_RUMAH_ENC'] = df_processed['STATUS_RUMAH_ENC'].fillna(0).astype(int)
    
    df_processed['PENDAPATAN_ZSCORE'] = (
        df_processed['PENDAPATAN ORANG TUA'] - scaler.mean_[0]
    ) / scaler.scale_[0]
    
    feature_cols = ['KELAS_NUM', 'PENDAPATAN_ZSCORE', 'PEKERJAAN_ENC', 
                   'JUMLAH TANGGUNGAN', 'STATUS_RUMAH_ENC']
    
    X = df_processed[feature_cols].values
    
    return df_input, X

def predict_batch(model, X, df_input):
    """Melakukan prediksi batch dan menambahkan hasil ke DataFrame."""
    predictions = model.predict(X)
    probabilities = model.predict_proba(X)
    
    df_result = df_input.copy()
    df_result['PREDIKSI'] = ['Penerima' if p == 1 else 'Tidak Penerima' for p in predictions]
    df_result['PROB_TIDAK (%)'] = (probabilities[:, 0] * 100).round(2)
    df_result['PROB_YA (%)'] = (probabilities[:, 1] * 100).round(2)
    
    return df_result

def evaluate_with_labeled_data(model, X, y_true, feature_cols):
    """Evaluasi model dengan data berlabel."""
    y_pred = model.predict(X)
    y_pred_proba = model.predict_proba(X)
    
    accuracy = accuracy_score(y_true, y_pred)
    cm = confusion_matrix(y_true, y_pred)
    
    TP = cm[1, 1] if cm.shape == (2, 2) else 0
    TN = cm[0, 0] if cm.shape == (2, 2) else 0
    FP = cm[0, 1] if cm.shape == (2, 2) else 0
    FN = cm[1, 0] if cm.shape == (2, 2) else 0
    
    prec0 = TN / (TN + FN) if (TN + FN) > 0 else 0
    rec0 = TN / (TN + FP) if (TN + FP) > 0 else 0
    f1_0 = 2 * prec0 * rec0 / (prec0 + rec0) if (prec0 + rec0) > 0 else 0
    
    prec1 = TP / (TP + FP) if (TP + FP) > 0 else 0
    rec1 = TP / (TP + FN) if (TP + FN) > 0 else 0
    f1_1 = 2 * prec1 * rec1 / (prec1 + rec1) if (prec1 + rec1) > 0 else 0
    
    sup0 = int((y_true == 0).sum())
    sup1 = int((y_true == 1).sum())
    tot = sup0 + sup1
    
    macro_prec = (prec0 + prec1) / 2
    macro_rec = (rec0 + rec1) / 2
    macro_f1 = (f1_0 + f1_1) / 2
    w_prec = (prec0 * sup0 + prec1 * sup1) / tot if tot > 0 else 0
    w_rec = (rec0 * sup0 + rec1 * sup1) / tot if tot > 0 else 0
    w_f1 = (f1_0 * sup0 + f1_1 * sup1) / tot if tot > 0 else 0
    
    return {
        'accuracy': accuracy,
        'cm': cm,
        'prec0': prec0, 'rec0': rec0, 'f1_0': f1_0,
        'prec1': prec1, 'rec1': rec1, 'f1_1': f1_1,
        'macro_prec': macro_prec, 'macro_rec': macro_rec, 'macro_f1': macro_f1,
        'w_prec': w_prec, 'w_rec': w_rec, 'w_f1': w_f1,
        'sup0': sup0, 'sup1': sup1,
        'y_pred': y_pred,
        'y_pred_proba': y_pred_proba
    }

def export_batch_prediction_excel(df_result: pd.DataFrame) -> bytes:
    """Export hasil prediksi batch ke Excel dengan styling profesional."""
    wb = Workbook()
    
    ws = wb.active
    ws.title = "Hasil Prediksi"
    
    ws.merge_cells('A1:H1')
    tc = ws['A1']
    tc.value = "HASIL PREDIKSI PENERIMA BSM - NAIVE BAYES"
    tc.font = Font(name='Inter', bold=True, size=14, color='FFFFFF')
    tc.fill = PatternFill("solid", start_color="065f46")
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells('A2:H2')
    sc = ws['A2']
    sc.value = f"Total Data: {len(df_result)} | Model: Gaussian Naive Bayes | Tanggal: {pd.Timestamp.now().strftime('%d-%m-%Y %H:%M')}"
    sc.font = Font(name='Inter', size=9, italic=True, color='595959')
    sc.fill = PatternFill("solid", start_color="d1fae5")
    sc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18
    
    headers = list(df_result.columns)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        _style_header(cell)
    ws.row_dimensions[3].height = 32
    
    row_colors = ["FFFFFF", "ecfdf5"]
    for ri, (_, row) in enumerate(df_result.iterrows(), 4):
        bg = row_colors[(ri - 4) % 2]
        for ci, val in enumerate(row, 1):
            safe_val = _safe_excel_value(val)
            cell = ws.cell(row=ri, column=ci, value=safe_val)
            
            col_name = headers[ci-1] if ci-1 < len(headers) else ""
            if col_name in ['PROB_TIDAK (%)', 'PROB_YA (%)']:
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                _style_data(cell, bg, align='right')
            elif col_name == 'PREDIKSI':
                if val == 'Penerima':
                    cell.fill = PatternFill("solid", start_color="d1fae5")
                    cell.font = Font(name='Inter', size=10, bold=True, color='065f46')
                else:
                    cell.fill = PatternFill("solid", start_color="fecaca")
                    cell.font = Font(name='Inter', size=10, bold=True, color='991b1b')
                cell.alignment = Alignment(horizontal='center', vertical='center')
                thin = Side(style='thin', color='D9D9D9')
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            elif col_name == 'PENDAPATAN ORANG TUA':
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                _style_data(cell, bg, align='right')
            else:
                _style_data(cell, bg)
    
    for i, header in enumerate(headers, 1):
        col_letter = get_column_letter(i)
        max_length = len(str(header))
        if not df_result.empty:
            col_values = df_result[header].astype(str)
            if len(col_values) > 0:
                max_val_len = col_values.str.len().max()
                max_length = max(max_length, max_val_len)
        adjusted_width = min(max_length + 4, 35)
        ws.column_dimensions[col_letter].width = adjusted_width
    
    ws2 = wb.create_sheet("Ringkasan")
    ws2.merge_cells('A1:B1')
    ws2['A1'].value = "RINGKASAN PREDIKSI"
    ws2['A1'].font = Font(name='Inter', bold=True, size=13, color='FFFFFF')
    ws2['A1'].fill = PatternFill("solid", start_color="065f46")
    ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[1].height = 26
    
    penerima = (df_result['PREDIKSI'] == 'Penerima').sum()
    tidak_penerima = (df_result['PREDIKSI'] == 'Tidak Penerima').sum()
    total = len(df_result)
    
    ringkasan_data = [
        ('Total Data', total),
        ('Jumlah Penerima', int(penerima)),
        ('Jumlah Tidak Penerima', int(tidak_penerima)),
        ('Persentase Penerima', f"{penerima/total*100:.1f}%" if total > 0 else "0%"),
        ('Persentase Tidak Penerima', f"{tidak_penerima/total*100:.1f}%" if total > 0 else "0%"),
    ]
    
    for i, (label, value) in enumerate(ringkasan_data, 3):
        cell_label = ws2.cell(row=i, column=1, value=label)
        _style_data(cell_label, align='left')
        cell_label.font = Font(name='Inter', size=10, bold=True)
        
        cell_value = ws2.cell(row=i, column=2, value=value)
        _style_data(cell_value)
    
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 22
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def handle_training_only_mode():
    """Menangani mode training only (testing 0%)."""
    st.markdown("---")
    st.markdown("### 📂 Upload Data Testing Manual")
    st.info("""
    ℹ️ Karena **Testing = 0%**, semua data digunakan untuk training dan tidak ada data testing otomatis.
    
    Silakan upload file Excel **berlabel** (memiliki kolom LABEL) untuk mengevaluasi model.
    File akan diproses menggunakan encoder dan scaler yang sama dari data training.
    """)
    
    uploaded_test = st.file_uploader(
        "Upload Data Testing Manual (.xlsx)",
        type=['xlsx'],
        key='test_upload_training_only',
        help="File Excel dengan kolom lengkap termasuk LABEL (Ya/Tidak)"
    )
    
    if uploaded_test is not None:
        try:
            df_test_raw = pd.read_excel(uploaded_test, sheet_name=0, header=1)
            if len(df_test_raw.columns) < 14:
                uploaded_test.seek(0)
                df_test_raw = pd.read_excel(uploaded_test, sheet_name=0)
            df_test_raw = df_test_raw.dropna(how='all')
            
            st.success(f"✅ File berhasil dimuat: {len(df_test_raw)} baris")
            
            le_pekerjaan = st.session_state.preprocess_info['le_pekerjaan']
            scaler = st.session_state.preprocess_info['scaler']
            
            df_test = df_test_raw.copy()
            
            df_test['PENDAPATAN ORANG TUA'] = pd.to_numeric(df_test['PENDAPATAN ORANG TUA'], errors='coerce')
            df_test['JUMLAH TANGGUNGAN'] = pd.to_numeric(df_test['JUMLAH TANGGUNGAN'], errors='coerce')
            df_test['PENDAPATAN ORANG TUA'] = df_test['PENDAPATAN ORANG TUA'].fillna(df_test['PENDAPATAN ORANG TUA'].mean())
            df_test['JUMLAH TANGGUNGAN'] = df_test['JUMLAH TANGGUNGAN'].fillna(df_test['JUMLAH TANGGUNGAN'].median())
            
            df_test['KELAS_NUM'] = df_test['KELAS'].astype(str).str.extract(r'(\d+)')
            df_test['KELAS_NUM'] = pd.to_numeric(df_test['KELAS_NUM'], errors='coerce').fillna(0).astype(int)
            
            pekerjaan_valid = set(le_pekerjaan.classes_)
            pekerjaan_input = set(df_test['PEKERJAAN ORANG TUA'].astype(str).unique())
            pekerjaan_unknown = pekerjaan_input - pekerjaan_valid
            
            if pekerjaan_unknown:
                st.error(f"❌ Kategori pekerjaan tidak dikenal: {', '.join(pekerjaan_unknown)}")
                st.info(f"Kategori yang valid: {', '.join(sorted(pekerjaan_valid))}")
                return
            
            df_test['PEKERJAAN ORANG TUA_ENC'] = le_pekerjaan.transform(
                df_test['PEKERJAAN ORANG TUA'].astype(str)
            )
            
            status_rumah_mapping = {"Milik Sendiri": 1, "Kontrak/sewa": 0}
            df_test['STATUS RUMAH_ENC'] = df_test['STATUS RUMAH'].map(status_rumah_mapping)
            df_test['STATUS RUMAH_ENC'] = df_test['STATUS RUMAH_ENC'].fillna(0).astype(int)
            
            df_test['PENDAPATAN_ZSCORE'] = scaler.transform(df_test[['PENDAPATAN ORANG TUA']])
            
            df_test['LABEL'] = df_test['LABEL'].astype(str).str.strip().str.capitalize()
            label_mapping = {"Ya": 1, "Tidak": 0}
            df_test['LABEL_ENC'] = df_test['LABEL'].map(label_mapping)
            
            valid_idx = df_test['LABEL_ENC'].notna()
            if valid_idx.sum() == 0:
                st.error("❌ Tidak ada label valid (Ya/Tidak) dalam data testing.")
                return
            
            df_test = df_test[valid_idx].copy()
            
            feature_cols = ['KELAS_NUM', 'PENDAPATAN_ZSCORE', 'PEKERJAAN ORANG TUA_ENC', 
                           'JUMLAH TANGGUNGAN', 'STATUS RUMAH_ENC']
            
            X_test = df_test[feature_cols].values
            y_test = df_test['LABEL_ENC'].values
            
            model = st.session_state.train_results['model']
            eval_results = evaluate_with_labeled_data(model, X_test, y_test, feature_cols)
            
            st.success(f"✅ Evaluasi selesai! Data testing: {len(df_test)} baris")
            
            col1, col2, col3, col4 = st.columns(4)
            with col1:
                st.metric("Akurasi Testing", f"{eval_results['accuracy']*100:.2f}%")
            with col2:
                st.metric("F1-Score (Macro)", f"{eval_results['macro_f1']*100:.2f}%")
            with col3:
                benar = (y_test == eval_results['y_pred']).sum()
                st.metric("Prediksi Benar", int(benar))
            with col4:
                salah = (y_test != eval_results['y_pred']).sum()
                st.metric("Prediksi Salah", int(salah))
            
            st.markdown("#### 📊 Confusion Matrix")
            cm = eval_results['cm']
            th = make_plotly_theme()
            colors = make_plotly_colors()
            fig = go.Figure(data=go.Heatmap(
                z=cm,
                x=['Predicted: Tidak', 'Predicted: Ya'],
                y=['Actual: Tidak', 'Actual: Ya'],
                text=[[str(v) for v in row] for row in cm],
                texttemplate="%{text}",
                textfont={"size": 22, "family": "JetBrains Mono", "color": "#FFFFFF"},
                colorscale=[[0, "#1E293B"], [0.5, "#3B82F6"], [1, "#93C5FD"]],
                showscale=True,
                colorbar=dict(title="Jumlah", tickfont=dict(color="#FFFFFF"), title_font=dict(color="#FFFFFF"))
            ))
            fig.update_layout(**th, height=380)
            st.plotly_chart(fig, use_container_width=True)
            
            with st.expander("📋 Detail Metrik Evaluasi", expanded=False):
                col_a, col_b = st.columns(2)
                with col_a:
                    st.markdown("**Per Kelas**")
                    st.dataframe(pd.DataFrame({
                        'Metrik': ['Precision', 'Recall', 'F1-Score'],
                        'Tidak Penerima (0)': [
                            f"{eval_results['prec0']:.4f}",
                            f"{eval_results['rec0']:.4f}",
                            f"{eval_results['f1_0']:.4f}"
                        ],
                        'Penerima (1)': [
                            f"{eval_results['prec1']:.4f}",
                            f"{eval_results['rec1']:.4f}",
                            f"{eval_results['f1_1']:.4f}"
                        ],
                    }), use_container_width=True, hide_index=True)
                with col_b:
                    st.markdown("**Average**")
                    st.dataframe(pd.DataFrame({
                        'Tipe': ['Macro', 'Weighted'],
                        'Precision': [f"{eval_results['macro_prec']:.4f}", f"{eval_results['w_prec']:.4f}"],
                        'Recall': [f"{eval_results['macro_rec']:.4f}", f"{eval_results['w_rec']:.4f}"],
                        'F1-Score': [f"{eval_results['macro_f1']:.4f}", f"{eval_results['w_f1']:.4f}"],
                    }), use_container_width=True, hide_index=True)
            
        except Exception as e:
            st.error(f"❌ Error: {e}")

def handle_testing_only_mode():
    """Menangani mode testing only (testing 100%)."""
    st.markdown("---")
    st.markdown("### 📂 Upload Data Training Manual")
    st.warning("""
    ⚠️ Karena **Testing = 100%**, semua data dianggap sebagai data testing dan tidak ada data training.
    
    Silakan upload file Excel untuk digunakan sebagai **data training**. 
    Setelah model dilatih, dataset awal akan digunakan sebagai data testing untuk evaluasi.
    """)
    
    uploaded_train = st.file_uploader(
        "Upload Data Training Manual (.xlsx)",
        type=['xlsx'],
        key='train_upload_testing_only',
        help="File Excel dengan kolom lengkap termasuk LABEL (Ya/Tidak)"
    )
    
    if uploaded_train is not None:
        try:
            df_train_raw = pd.read_excel(uploaded_train, sheet_name=0, header=1)
            if len(df_train_raw.columns) < 14:
                uploaded_train.seek(0)
                df_train_raw = pd.read_excel(uploaded_train, sheet_name=0)
            df_train_raw = df_train_raw.dropna(how='all')
            
            st.info(f"Memproses {len(df_train_raw)} baris data training...")
            
            df_train_p, info = run_preprocessing(df_train_raw)
            
            st.success(f"✅ Data training berhasil diproses: {info['total_rows']} baris")
            
            if st.session_state.df_processed is not None:
                df_test = st.session_state.df_processed
                
                feature_cols = ['KELAS_NUM', 'PENDAPATAN_ZSCORE', 'PEKERJAAN ORANG TUA_ENC', 
                               'JUMLAH TANGGUNGAN', 'STATUS RUMAH_ENC']
                
                X_train = df_train_p[feature_cols].values
                y_train = df_train_p['LABEL_ENC'].values
                
                valid_test = df_test[feature_cols].notna().all(axis=1) & df_test['LABEL_ENC'].notna()
                X_test = df_test[feature_cols][valid_test].values
                y_test = df_test['LABEL_ENC'][valid_test].values
                
                if len(X_test) == 0:
                    st.error("❌ Tidak ada data testing valid.")
                    return
                
                model = GaussianNB()
                model.fit(X_train, y_train)
                
                eval_results = evaluate_with_labeled_data(model, X_test, y_test, feature_cols)
                
                st.success(f"✅ Training dan evaluasi selesai!")
                
                col1, col2, col3, col4 = st.columns(4)
                with col1:
                    st.metric("Akurasi Testing", f"{eval_results['accuracy']*100:.2f}%")
                with col2:
                    st.metric("F1-Score (Macro)", f"{eval_results['macro_f1']*100:.2f}%")
                with col3:
                    benar = (y_test == eval_results['y_pred']).sum()
                    st.metric("Prediksi Benar", int(benar))
                with col4:
                    salah = (y_test != eval_results['y_pred']).sum()
                    st.metric("Prediksi Salah", int(salah))
                
                st.markdown("#### 📊 Confusion Matrix")
                cm = eval_results['cm']
                th = make_plotly_theme()
                colors = make_plotly_colors()
                fig = go.Figure(data=go.Heatmap(
                    z=cm,
                    x=['Predicted: Tidak', 'Predicted: Ya'],
                    y=['Actual: Tidak', 'Actual: Ya'],
                    text=[[str(v) for v in row] for row in cm],
                    texttemplate="%{text}",
                    textfont={"size": 22, "family": "JetBrains Mono", "color": "#FFFFFF"},
                    colorscale=[[0, "#1E293B"], [0.5, "#3B82F6"], [1, "#93C5FD"]],
                    showscale=True,
                    colorbar=dict(title="Jumlah", tickfont=dict(color="#FFFFFF"), title_font=dict(color="#FFFFFF"))
                ))
                fig.update_layout(**th, height=380)
                st.plotly_chart(fig, use_container_width=True)
                
                with st.expander("📋 Detail Metrik Evaluasi", expanded=False):
                    col_a, col_b = st.columns(2)
                    with col_a:
                        st.markdown("**Per Kelas**")
                        st.dataframe(pd.DataFrame({
                            'Metrik': ['Precision', 'Recall', 'F1-Score'],
                            'Tidak Penerima (0)': [
                                f"{eval_results['prec0']:.4f}",
                                f"{eval_results['rec0']:.4f}",
                                f"{eval_results['f1_0']:.4f}"
                            ],
                            'Penerima (1)': [
                                f"{eval_results['prec1']:.4f}",
                                f"{eval_results['rec1']:.4f}",
                                f"{eval_results['f1_1']:.4f}"
                            ],
                        }), use_container_width=True, hide_index=True)
                    with col_b:
                        st.markdown("**Average**")
                        st.dataframe(pd.DataFrame({
                            'Tipe': ['Macro', 'Weighted'],
                            'Precision': [f"{eval_results['macro_prec']:.4f}", f"{eval_results['w_prec']:.4f}"],
                            'Recall': [f"{eval_results['macro_rec']:.4f}", f"{eval_results['w_rec']:.4f}"],
                            'F1-Score': [f"{eval_results['macro_f1']:.4f}", f"{eval_results['w_f1']:.4f}"],
                        }), use_container_width=True, hide_index=True)
                
        except Exception as e:
            st.error(f"❌ Error: {e}")

def show_prediction_page():
    """Halaman untuk prediksi data baru (manual dan batch)."""
    render_sticky_header("🔮", "Prediksi Data Baru")
    
    if st.session_state.train_results is None:
        st.warning("⚠️ Model belum dilatih. Silakan lakukan training model terlebih dahulu di menu **🤖 Training Model**.")
        return
    
    if st.session_state.preprocess_info is None:
        st.warning("⚠️ Data preprocessing belum tersedia. Silakan upload dan preprocessing data terlebih dahulu di menu **📂 Data & Preprocessing**.")
        return
    
    model = st.session_state.train_results['model']
    
    if model is None:
        st.warning("⚠️ Model tidak tersedia (Testing 100%). Silakan upload data training manual di Dashboard.")
        return
    
    tabs = st.tabs(["✍️ Input Manual", "📂 Prediksi Massal"])
    
    with tabs[0]:
        st.markdown("### ✍️ Prediksi Manual")
        st.markdown("Masukkan data siswa secara manual untuk memprediksi status penerimaan BSM.")
        
        with st.form("prediction_manual_form"):
            col1, col2 = st.columns(2)
            
            with col1:
                kelas = st.selectbox("📚 Kelas", options=["7", "8", "9"])
                pendapatan = st.number_input("💰 Pendapatan Orang Tua (Rp)", min_value=0, value=2000000, step=100000, format="%d")
                pekerjaan_options = list(st.session_state.preprocess_info['le_pekerjaan'].classes_)
                pekerjaan = st.selectbox("💼 Pekerjaan Orang Tua", options=pekerjaan_options)
            
            with col2:
                tanggungan = st.number_input("👨‍👩‍👧‍👦 Jumlah Tanggungan", min_value=0, max_value=20, value=3, step=1)
                status_rumah = st.selectbox("🏠 Status Rumah", options=["Milik Sendiri", "Kontrak/sewa"])
            
            submitted = st.form_submit_button("🔮 Prediksi Sekarang", use_container_width=True, type="primary")
        
        if submitted:
            with st.spinner("🔮 Melakukan prediksi..."):
                time.sleep(0.5)
                
                features = preprocess_single_input(
                    kelas=kelas, pendapatan=pendapatan, pekerjaan=pekerjaan,
                    tanggungan=tanggungan, status_rumah=status_rumah
                )
                
                prediction, probabilities = predict_single(model, features)
                
                st.markdown("---")
                st.markdown("### 🎯 Hasil Prediksi")
                
                if prediction == 1:
                    card_class = "success"
                    icon = "✅"
                    result_text = "PENERIMA BSM"
                    badge_class = "badge-penerima"
                else:
                    card_class = "danger"
                    icon = "❌"
                    result_text = "TIDAK PENERIMA BSM"
                    badge_class = "badge-tidak"
                
                st.markdown(f"""
                <div class="prediction-card {card_class}">
                    <div class="prediction-icon">{icon}</div>
                    <div class="prediction-result">{result_text}</div>
                    <div style="margin: 1rem 0;">
                        <span class="{badge_class}">{result_text}</span>
                    </div>
                </div>
                """, unsafe_allow_html=True)
                
                col_prob1, col_prob2 = st.columns(2)
                with col_prob1:
                    st.metric("Probabilitas TIDAK", f"{probabilities[0]*100:.2f}%")
                with col_prob2:
                    st.metric("Probabilitas YA", f"{probabilities[1]*100:.2f}%")
                
                with st.expander("📋 Detail Input & Perhitungan", expanded=False):
                    col_a, col_b = st.columns(2)
                    with col_a:
                        st.markdown("**Data Input:**")
                        st.markdown(f"- 📚 Kelas: **{kelas}**")
                        st.markdown(f"- 💰 Pendapatan: **Rp {pendapatan:,.0f}**")
                        st.markdown(f"- 💼 Pekerjaan: **{pekerjaan}**")
                    with col_b:
                        st.markdown("**Hasil Encoding & Standardisasi:**")
                        st.markdown(f"- 👨‍👩‍👧‍👦 Tanggungan: **{tanggungan}**")
                        st.markdown(f"- 🏠 Status Rumah: **{status_rumah}**")
                        st.markdown(f"- 📐 Z-Score Pendapatan: **{features[0][1]:.6f}**")
    
    with tabs[1]:
        st.markdown("### 📂 Prediksi Massal")
        st.markdown("Upload file Excel berisi data siswa untuk diprediksi secara massal sekaligus.")
        
        with st.expander("📌 Panduan Format File untuk Prediksi Massal", expanded=False):
            st.markdown("""
            **Format file Excel (.xlsx) yang wajib dipenuhi:**
            
            | Kolom | Tipe Data | Keterangan | Contoh |
            |-------|-----------|------------|--------|
            | KELAS | string/angka | Tingkat kelas siswa | 7, 8, atau 9 |
            | PENDAPATAN ORANG TUA | angka | Pendapatan bulanan (Rp) | 2000000 |
            | PEKERJAAN ORANG TUA | string | Jenis pekerjaan | Buruh, Petani, dll |
            | JUMLAH TANGGUNGAN | angka | Jumlah tanggungan keluarga | 3 |
            | STATUS RUMAH | string | Status kepemilikan | "Milik Sendiri" / "Kontrak/sewa" |
            
            **⚠️ Penting:**
            - Kategori **PEKERJAAN ORANG TUA** harus sesuai dengan yang ada di data training.
            - File harus dalam format **.xlsx** (bukan .xls).
            """)
        
        uploaded_file = st.file_uploader(
            "Upload File Excel untuk Prediksi Massal",
            type=['xlsx'],
            help="Upload file .xlsx dengan kolom sesuai panduan di atas"
        )
        
        if uploaded_file is not None:
            try:
                df_input = pd.read_excel(uploaded_file)
                st.success(f"✅ File berhasil dimuat: **{len(df_input)}** baris × **{len(df_input.columns)}** kolom")
                
                st.markdown("#### 📋 Preview Data Input (10 baris pertama)")
                st.dataframe(df_input.head(10), use_container_width=True)
                st.caption(f"Kolom tersedia: {', '.join(df_input.columns.tolist())}")
                
                if st.button("📊 Jalankan Prediksi Massal", use_container_width=True, type="primary", key="btn_batch_predict"):
                    with st.spinner("🔄 Memproses data dan melakukan prediksi..."):
                        progress_bar = st.progress(0)
                        for i in range(1, 101, 20):
                            time.sleep(0.02)
                            progress_bar.progress(i)
                        
                        try:
                            df_input_copy, X = preprocess_batch_data(df_input)
                            df_result = predict_batch(model, X, df_input_copy)
                            st.session_state.batch_prediction_df = df_result
                            progress_bar.progress(100)
                            st.success("✅ Prediksi massal selesai!")
                            time.sleep(0.5)
                            st.rerun()
                        except ValueError as ve:
                            st.error(f"❌ Error Validasi: {ve}")
                            progress_bar.progress(100)
                        except Exception as e:
                            st.error(f"❌ Error: {e}")
                
                if st.session_state.batch_prediction_df is not None:
                    df_result = st.session_state.batch_prediction_df
                    
                    st.markdown("---")
                    st.markdown("### 📊 Ringkasan Hasil Prediksi")
                    
                    penerima = int((df_result['PREDIKSI'] == 'Penerima').sum())
                    tidak_penerima = int((df_result['PREDIKSI'] == 'Tidak Penerima').sum())
                    total = len(df_result)
                    
                    col_a, col_b, col_c = st.columns(3)
                    with col_a:
                        st.metric("📦 Total Data", total)
                    with col_b:
                        pct_penerima = penerima / total * 100 if total > 0 else 0
                        st.metric("✅ Penerima BSM", penerima, delta=f"{pct_penerima:.1f}%")
                    with col_c:
                        pct_tidak = tidak_penerima / total * 100 if total > 0 else 0
                        st.metric("❌ Tidak Penerima", tidak_penerima, delta=f"{pct_tidak:.1f}%", delta_color="inverse")
                    
                    st.markdown("#### 📋 Tabel Hasil Prediksi")
                    st.dataframe(
                        df_result, use_container_width=True, height=450, hide_index=True,
                        column_config={
                            'PROB_TIDAK (%)': st.column_config.NumberColumn(format="%.2f%%"),
                            'PROB_YA (%)': st.column_config.NumberColumn(format="%.2f%%"),
                        }
                    )
                    
                    st.markdown("---")
                    excel_bytes = export_batch_prediction_excel(df_result)
                    
                    st.download_button(
                        label="📥 Download Hasil Prediksi (.xlsx)",
                        data=excel_bytes,
                        file_name="hasil_prediksi_bsm.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )
                    
            except Exception as e:
                st.error(f"❌ Error membaca file: {e}")

# ─────────────────────────────────────────────────────────────
# LANDING PAGE
# ─────────────────────────────────────────────────────────────

def show_landing():
    """Halaman awal sebelum masuk dashboard."""
    
    st.markdown("""
    <div class="hero-home">
        <div style="position:absolute;top:18px;right:18px;width:60px;height:60px;
                    background:#FFE500;border:3px solid #000;transform:rotate(12deg);
                    box-shadow:3px 3px 0 #000;"></div>
        <div style="position:absolute;bottom:24px;right:80px;width:40px;height:40px;
                    border:3px solid #FFE500;border-radius:50%;"></div>
        <div style="position:absolute;top:50%;left:calc(100% - 160px);
                    font-family:'Bungee',cursive;font-size:6rem;color:rgba(255,255,255,0.06);
                    line-height:1;pointer-events:none;white-space:nowrap;">BSM</div>
        <div class="hero-badge">
            <span>🎓</span> MACHINE LEARNING &nbsp;·&nbsp; NAIVE BAYES &nbsp;·&nbsp; 2024
        </div>
        <div class="hero-title">Sistem Klasifikasi<br>Penerima BSM</div>
        <div class="hero-sub">
            Dashboard cerdas berbasis <strong>Gaussian Naive Bayes</strong> — membantu pengambilan 
            keputusan pemberian bantuan secara objektif, transparan, dan terukur.
        </div>
        <div style="display:flex; gap:12px; flex-wrap:wrap; margin-bottom:2rem;">
            <div style="background:#FFE500;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;">5</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Fitur Input</div>
            </div>
            <div style="background:#00F5FF;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;">2</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Kelas Output</div>
            </div>
            <div style="background:#B6FF00;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;">GNB</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Algoritma</div>
            </div>
            <div style="background:#fff;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;color:#2146FF;">100%</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Open Source</div>
            </div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns(3)
    with col1:
        render_metric_card("🎯", "Objektif & Adil", "Probabilistik", "Klasifikasi berbasis probabilitas tanpa bias subjektif")
    with col2:
        render_metric_card("⚡", "Cepat & Ringan", "Real-time", "Training model dalam hitungan milidetik")
    with col3:
        render_metric_card("📊", "Transparan", "Terukur", "Semua metrik evaluasi divisualisasikan dengan jelas")

    st.markdown("<br>", unsafe_allow_html=True)

    st.markdown("### 📘 Penjelasan Metode & Rumus")
    with st.expander("🔍 Klik untuk melihat teori Gaussian Naive Bayes dan rumus lengkap", expanded=False):
        st.markdown("""
        **Gaussian Naive Bayes (GNB)** adalah algoritma klasifikasi probabilistik yang mengasumsikan 
        setiap fitur mengikuti distribusi normal (Gaussian).
        """)
        
        col_a, col_b = st.columns(2)
        with col_a:
            st.markdown("#### 🧠 Teorema Bayes")
            st.latex(r"P(C_k|X) = \frac{P(X|C_k) \cdot P(C_k)}{P(X)}")
            st.markdown("#### 📈 Gaussian PDF")
            st.latex(r"P(x_i|C_k) = \frac{1}{\sqrt{2\pi\sigma_k^2}} \exp\left(-\frac{(x_i-\mu_k)^2}{2\sigma_k^2}\right)")
            st.markdown("#### 📐 Z-Score")
            st.latex(r"z = \frac{x - \mu}{\sigma}")
        with col_b:
            st.markdown("#### 🎯 Accuracy")
            st.latex(r"\text{Accuracy} = \frac{TP + TN}{TP + TN + FP + FN}")
            st.markdown("#### 📊 Precision")
            st.latex(r"\text{Precision} = \frac{TP}{TP + FP}")
            st.markdown("#### 🔍 Recall")
            st.latex(r"\text{Recall} = \frac{TP}{TP + FN}")
            st.markdown("#### ⚖️ F1-Score")
            st.latex(r"F_1 = 2 \cdot \frac{\text{Precision} \cdot \text{Recall}}{\text{Precision} + \text{Recall}}")
    
    st.markdown("<br>", unsafe_allow_html=True)
    
    st.markdown("### 🔄 Alur Kerja Sistem")
    steps = [
        ("📂", "Upload Dataset", "Unggah file Excel (.xlsx) data BSM"),
        ("⚙️", "Preprocessing", "Missing value, duplikasi, encoding, standardisasi"),
        ("📊", "Visualisasi", "Eksplorasi data dengan grafik Plotly"),
        ("🤖", "Training Model", "Latih Gaussian Naive Bayes"),
        ("📈", "Evaluasi", "Accuracy, Precision, Recall, F1-Score"),
        ("🔮", "Prediksi Baru", "Prediksi manual / batch dari file Excel"),
        ("💾", "Export", "Unduh hasil ke file Excel profesional"),
    ]
    for icon, title, desc in steps:
        st.markdown(f"""
        <div style="display:flex; align-items:center; gap:14px; padding:12px 16px;
                    background:#fff; border:3px solid #000; box-shadow:4px 4px 0 #000;
                    margin-bottom:8px; transition:transform 0.15s;">
            <div style="font-size:1.6rem; width:44px; text-align:center;
                        background:#FFE500; border:2px solid #000; padding:6px;
                        box-shadow:2px 2px 0 #000; flex-shrink:0;">{icon}</div>
            <div>
                <div style="font-family:'Archivo Black',sans-serif; font-weight:900; color:#000;
                            font-size:0.88rem; text-transform:uppercase; letter-spacing:0.04em;">{title}</div>
                <div style="font-family:'IBM Plex Mono',monospace; font-size:0.75rem; color:#555;
                            margin-top:2px;">{desc}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)
    
    st.markdown("<br>", unsafe_allow_html=True)
    
    _, center_col, _ = st.columns([1, 1.5, 1])
    with center_col:
        if st.button("🚀 Mulai Menggunakan Sistem", use_container_width=True, type="primary"):
            st.session_state.app_state = 'main'
            st.rerun()
    
    st.markdown("""
    <div class="footer">
        <div class="footer-brand">🎓 SISTEM KLASIFIKASI PENERIMA BSM</div>
        <div class="footer-divider"></div>
        <div style="font-family:'IBM Plex Mono',monospace;font-size:0.8rem;letter-spacing:0.06em;">
            GAUSSIAN NAIVE BAYES &nbsp;·&nbsp; MACHINE LEARNING &nbsp;·&nbsp; STREAMLIT
        </div>
        <div style="margin-top:0.5rem;font-family:'IBM Plex Mono',monospace;
                    color:rgba(255,255,255,0.4);font-size:0.7rem;letter-spacing:0.08em;">
            // DIBUAT UNTUK KEPERLUAN AKADEMIK DAN PENELITIAN //
        </div>
    </div>
    """, unsafe_allow_html=True)

# ─────────────────────────────────────────────────────────────
# MAIN APPLICATION
# ─────────────────────────────────────────────────────────────

def show_main():
    """Dashboard utama dengan sidebar navigasi."""
    
    with st.sidebar:
        st.markdown("""
        <div style="text-align:center; padding:1.4rem 0 1rem;">
            <div style="font-size:2.8rem; filter: drop-shadow(3px 3px 0 #000);">🎓</div>
            <div style="font-family:'Bungee',cursive; font-size:1.1rem; margin-top:6px; color:#FFE500;
                        letter-spacing:0.08em; text-shadow: 2px 2px 0 #000;">KLASIFIKASI BSM</div>
            <div style="font-family:'IBM Plex Mono',monospace; font-size:0.65rem; color:rgba(255,255,255,0.7);
                        margin-top:4px; letter-spacing:0.12em; text-transform:uppercase;">// NAIVE BAYES //</div>
            <div style="width:60%; height:3px; background:#FFE500; margin:10px auto 0;
                        border:1px solid #000;"></div>
        </div>
        """, unsafe_allow_html=True)
        st.markdown("---")
        
        page = st.radio(
            "Navigasi",
            [
                "🏠 Dashboard",
                "📂 Data & Preprocessing",
                "📊 Visualisasi",
                "🤖 Training Model",
                "📈 Evaluasi",
                "🗂️ Hasil Prediksi",
                "🔮 Prediksi Data Baru",
                "💾 Export Hasil",
                "ℹ️ Tentang Sistem",
            ],
            label_visibility="collapsed"
        )
        
        st.markdown("---")
        st.markdown("**Status Sistem**")
        ds_loaded = st.session_state.df_processed is not None
        model_done = st.session_state.train_results is not None
        
        st.markdown(f"{'🟢' if ds_loaded else '⚪'} Dataset {'siap' if ds_loaded else 'belum dimuat'}")
        st.markdown(f"{'🟢' if model_done else '⚪'} Model {'terlatih' if model_done else 'belum dilatih'}")
        
        if ds_loaded:
            st.markdown("---")
            st.markdown("**Pengaturan Split Data**")
            
            if st.session_state.training_locked:
                test_size = st.slider(
                    "Ukuran Data Uji (%)", 0, 100, int(st.session_state.test_size * 100), 5,
                    disabled=True, help="Training sudah dikunci. Klik 'Reset Training' untuk mengubah."
                ) / 100.0
            else:
                test_size = st.slider(
                    "Ukuran Data Uji (%)", 0, 100, int(st.session_state.test_size * 100), 5,
                    help="Geser untuk mengatur proporsi data testing"
                ) / 100.0
                st.session_state.test_size = test_size
            
            if test_size == 0.0:
                st.info("ℹ️ Testing 0%: Semua data untuk training")
            elif test_size == 1.0:
                st.warning("⚠️ Testing 100%: Semua data untuk testing")
            else:
                st.info(f"ℹ️ Split: {(1-test_size)*100:.0f}% Training | {test_size*100:.0f}% Testing")
        
        st.markdown("---")
        if st.button("🏠 Kembali ke Halaman Awal", use_container_width=True):
            st.session_state.app_state = 'home'
            st.rerun()

    # ═════════════════════════════════════════════════════════
    # PAGE ROUTING
    # ═════════════════════════════════════════════════════════

    if page == "🏠 Dashboard":
        render_sticky_header("🏠", "Dashboard Utama")
        
        if st.session_state.df_processed is not None:
            df_p = st.session_state.df_processed
            n_total = len(df_p)
            n_ya = int((df_p['LABEL_ENC'] == 1).sum())
            n_tdk = int((df_p['LABEL_ENC'] == 0).sum())
            
            render_metric_cards([
                {"icon": "🗃️", "label": "Total Data", "value": str(n_total), "desc": "setelah preprocessing"},
                {"icon": "✅", "label": "Penerima BSM", "value": str(n_ya), "desc": f"{n_ya/n_total*100:.1f}% dari total"},
                {"icon": "❌", "label": "Tidak Penerima", "value": str(n_tdk), "desc": f"{n_tdk/n_total*100:.1f}% dari total"},
                {"icon": "📐", "label": "Jumlah Fitur", "value": "5", "desc": "fitur input model"},
            ])
            
            if st.session_state.train_results:
                res = st.session_state.train_results
                st.markdown("<br>", unsafe_allow_html=True)
                
                metrics_display = [
                    {"icon": "🎯", "label": "Akurasi Training", "value": f"{res['train_acc']*100:.1f}%" if res['train_acc'] is not None else "N/A", "desc": "data latih"},
                    {"icon": "🔬", "label": "Algoritma", "value": "GNB", "desc": "Gaussian Naive Bayes"},
                ]
                
                if res['test_acc'] is not None:
                    metrics_display.insert(1, {"icon": "🏁", "label": "Akurasi Testing", "value": f"{res['test_acc']*100:.1f}%", "desc": "data uji"})
                else:
                    metrics_display.insert(1, {"icon": "🏁", "label": "Akurasi Testing", "value": "N/A", "desc": "testing 0%"})
                
                render_metric_cards(metrics_display)
            
            test_size = st.session_state.test_size
            if test_size == 0.0 and st.session_state.train_results is not None:
                handle_training_only_mode()
            elif test_size == 1.0:
                handle_testing_only_mode()
        else:
            st.info("👋 Selamat datang! Silakan menuju menu **📂 Data & Preprocessing** untuk memulai.")

    elif page == "📂 Data & Preprocessing":
        render_sticky_header("📂", "Upload & Preprocessing Dataset")
        
        uploaded = st.file_uploader(
            "Upload file Excel (.xlsx) data BSM",
            type=["xlsx"],
            help="Format: .xlsx | Header di baris ke-2 | Kolom sesuai template BSM"
        )
        
        if uploaded is not None:
            try:
                df_raw = pd.read_excel(uploaded, sheet_name=0, header=1)
                if len(df_raw.columns) < 14:
                    uploaded.seek(0)
                    df_raw = pd.read_excel(uploaded, sheet_name=0)
                df_raw = df_raw.dropna(how='all')
                st.session_state.df_raw = df_raw
                
                st.success(f"✅ File berhasil dimuat: **{len(df_raw)}** baris × **{len(df_raw.columns)}** kolom")
                
                st.markdown("### 📋 Preview Dataset Lengkap")
                st.dataframe(df_raw, use_container_width=True, height=500)
                
                st.markdown("---")
                
                col_btn, col_info = st.columns([1, 3])
                with col_btn:
                    if st.button("⚙️ Jalankan Preprocessing", use_container_width=True, type="primary"):
                        with st.spinner("🔄 Memproses data..."):
                            progress_bar = st.progress(0)
                            for i in range(1, 101):
                                time.sleep(0.005)
                                progress_bar.progress(i)
                            df_p, info = run_preprocessing(df_raw)
                            st.session_state.df_processed = df_p
                            st.session_state.preprocess_info = info
                            st.session_state.training_locked = False
                            st.session_state.train_results = None
                            st.session_state.batch_prediction_df = None
                        st.rerun()
                
                if st.session_state.df_processed is not None:
                    st.markdown("---")
                    st.markdown("### ✅ Hasil Preprocessing")
                    
                    info = st.session_state.preprocess_info
                    col_a, col_b, col_c = st.columns(3)
                    with col_a:
                        st.metric("Data Awal", info['row_after_dup'] + info['dup_removed'])
                    with col_b:
                        st.metric("Duplikat Dihapus", info['dup_removed'])
                    with col_c:
                        st.metric("Data Final", info['row_after_dup'])
                    
                    st.dataframe(st.session_state.df_processed, use_container_width=True, height=500)
                    
            except Exception as e:
                st.error(f"❌ Error membaca file: {e}")
        
        with st.expander("📌 Panduan Format File", expanded=False):
            st.markdown("""
            | Kolom | Tipe Data | Deskripsi |
            |-------|-----------|-----------|
            | NO | integer | Nomor urut data |
            | NAMA PESERTA DIDIK | string | Nama lengkap siswa |
            | SEKOLAH | string | Nama sekolah |
            | KELAS | string | Tingkat kelas (7, 8, atau 9) |
            | PENDAPATAN ORANG TUA | float | Pendapatan orang tua (Rp) |
            | PEKERJAAN ORANG TUA | string | Jenis pekerjaan orang tua |
            | JUMLAH TANGGUNGAN | integer | Jumlah tanggungan keluarga |
            | STATUS RUMAH | string | Milik Sendiri / Kontrak/sewa |
            | LABEL | string | Ya / Tidak |
            """)

    elif page == "📊 Visualisasi":
        render_sticky_header("📊", "Visualisasi Data")
        
        if st.session_state.df_processed is None:
            st.warning("⚠️ Dataset belum diproses.")
        else:
            df = st.session_state.df_processed
            th = make_plotly_theme()
            colors = make_plotly_colors()
            
            tabs = st.tabs(["🥧 Distribusi Label", "💰 Pendapatan", "💼 Pekerjaan", "🏠 Status Rumah", "🎓 Kelas", "📊 Korelasi"])
            
            with tabs[0]:
                col1, col2 = st.columns([1, 1])
                with col1:
                    label_counts = df['LABEL'].value_counts()
                    fig = px.pie(
                        names=label_counts.index, 
                        values=label_counts.values,
                        color_discrete_sequence=[colors['penerima'], colors['tidak_penerima']], 
                        hole=0.45,
                        title="Distribusi Label Penerima BSM"
                    )
                    fig.update_traces(
                        textposition='outside', 
                        textinfo='label+percent+value', 
                        textfont_size=13,
                        textfont_color='#FFFFFF'
                    )
                    fig.update_layout(
                        **th, 
                        height=420,
                        legend=dict(font=dict(color='#FFFFFF'))
                    )
                    st.plotly_chart(fig, use_container_width=True)
                with col2:
                    st.markdown("<br>", unsafe_allow_html=True)
                    col_met1, col_met2 = st.columns(2)
                    with col_met1:
                        st.metric("Total Penerima", int((df['LABEL_ENC']==1).sum()))
                    with col_met2:
                        st.metric("Total Tidak Penerima", int((df['LABEL_ENC']==0).sum()))
                    st.metric("Rasio", f"{(df['LABEL_ENC']==1).sum()/len(df)*100:.1f}% Penerima")
            
            with tabs[1]:
                col1, col2 = st.columns(2)
                with col1:
                    fig1 = px.histogram(
                        df, x='PENDAPATAN ORANG TUA', nbins=30,
                        color_discrete_sequence=[colors['primary']],
                        title="Histogram Pendapatan (Rupiah)"
                    )
                    fig1.update_layout(**th, height=380)
                    st.plotly_chart(fig1, use_container_width=True)
                with col2:
                    fig2 = px.histogram(
                        df, x='PENDAPATAN_ZSCORE', nbins=30,
                        color_discrete_sequence=[colors['info']],
                        title="Histogram Pendapatan (Z-Score)"
                    )
                    fig2.update_layout(**th, height=380)
                    st.plotly_chart(fig2, use_container_width=True)
                
                fig3 = px.box(
                    df, x='LABEL', y='PENDAPATAN ORANG TUA', color='LABEL',
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    title="Box Plot Pendapatan per Status Penerima"
                )
                fig3.update_layout(**th, height=380)
                st.plotly_chart(fig3, use_container_width=True)
            
            with tabs[2]:
                pek_cnt = df['PEKERJAAN ORANG TUA'].value_counts().reset_index()
                pek_cnt.columns = ['Pekerjaan', 'Jumlah']
                fig = px.bar(
                    pek_cnt, y='Pekerjaan', x='Jumlah', orientation='h',
                    color='Jumlah', 
                    color_continuous_scale=colors['gradient_blue'],
                    title="Frekuensi Pekerjaan Orang Tua", 
                    text='Jumlah'
                )
                fig.update_traces(textposition='outside', textfont_color='#FFFFFF')
                fig.update_layout(**th, height=max(350, len(pek_cnt)*38),
                                 coloraxis_showscale=False, yaxis_categoryorder='total ascending')
                st.plotly_chart(fig, use_container_width=True)
                
                pek_label = df.groupby(['PEKERJAAN ORANG TUA', 'LABEL']).size().reset_index(name='Jumlah')
                fig2 = px.bar(
                    pek_label, y='PEKERJAAN ORANG TUA', x='Jumlah', color='LABEL',
                    orientation='h', 
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    barmode='stack', 
                    title="Pekerjaan vs Status Penerima"
                )
                fig2.update_layout(**th, height=max(350, len(pek_cnt)*38), yaxis_categoryorder='total ascending')
                st.plotly_chart(fig2, use_container_width=True)
            
            with tabs[3]:
                col1, col2 = st.columns(2)
                with col1:
                    sr_cnt = df['STATUS RUMAH'].value_counts()
                    fig1 = px.pie(
                        names=sr_cnt.index, values=sr_cnt.values,
                        color_discrete_sequence=[colors['primary'], colors['info']],
                        hole=0.4, 
                        title="Distribusi Status Rumah"
                    )
                    fig1.update_traces(textposition='inside', textinfo='label+percent', textfont_color='#FFFFFF')
                    fig1.update_layout(**th, height=350)
                    st.plotly_chart(fig1, use_container_width=True)
                with col2:
                    sr_label = df.groupby(['STATUS RUMAH', 'LABEL']).size().reset_index(name='Jumlah')
                    fig2 = px.bar(
                        sr_label, x='STATUS RUMAH', y='Jumlah', color='LABEL',
                        barmode='group', 
                        color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                        title="Status Rumah vs Status Penerima"
                    )
                    fig2.update_layout(**th, height=350)
                    st.plotly_chart(fig2, use_container_width=True)
            
            with tabs[4]:
                kl_cnt = df.groupby(['KELAS', 'LABEL']).size().reset_index(name='Jumlah')
                fig = px.bar(
                    kl_cnt, x='KELAS', y='Jumlah', color='LABEL', barmode='group',
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    title="Distribusi Kelas berdasarkan Status Penerima"
                )
                fig.update_layout(**th, height=400)
                st.plotly_chart(fig, use_container_width=True)
                
                fig2 = px.histogram(
                    df, x='JUMLAH TANGGUNGAN', color='LABEL',
                    barmode='overlay', opacity=0.75,
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    title="Distribusi Jumlah Tanggungan per Status Penerima"
                )
                fig2.update_layout(**th, height=380)
                st.plotly_chart(fig2, use_container_width=True)
            
            with tabs[5]:
                num_cols = ['KELAS_NUM','PENDAPATAN_ZSCORE','PEKERJAAN ORANG TUA_ENC',
                           'JUMLAH TANGGUNGAN','STATUS RUMAH_ENC','LABEL_ENC']
                num_avail = [c for c in num_cols if c in df.columns]
                corr = df[num_avail].corr().round(3)
                fig = px.imshow(
                    corr, text_auto=True, 
                    color_continuous_scale=colors['gradient_blue'],
                    title="Heatmap Korelasi Fitur", 
                    aspect='auto'
                )
                fig.update_layout(**th, height=500)
                st.plotly_chart(fig, use_container_width=True)

    elif page == "🤖 Training Model":
        render_sticky_header("🤖", "Training Model")
        
        if st.session_state.df_processed is None:
            st.warning("⚠️ Dataset belum diproses.")
        else:
            df = st.session_state.df_processed
            test_size = st.session_state.test_size
            n_valid = df['LABEL_ENC'].notna().sum()
            n_train = int(n_valid * (1 - test_size))
            n_test = n_valid - n_train
            
            render_metric_cards([
                {"icon": "📦", "label": "Total Data Valid", "value": str(n_valid)},
                {"icon": "🏋️", "label": "Data Training", "value": str(n_train), "desc": f"{(1-test_size)*100:.0f}%"},
                {"icon": "🧪", "label": "Data Testing", "value": str(n_test), "desc": f"{test_size*100:.0f}%"},
                {"icon": "🔬", "label": "Algoritma", "value": "GNB", "desc": "Gaussian Naive Bayes"},
            ])
            
            if test_size == 1.0:
                st.error("❌ Testing 100% tidak valid untuk training otomatis.")
                st.info("ℹ️ Silakan upload data training manual di Dashboard atau ubah split data.")
            elif test_size == 0.0:
                alert("⚠️ Testing 0%: Semua data untuk training. Evaluasi testing tidak tersedia otomatis.", "warning")
            
            st.markdown("---")
            
            col_btn1, col_btn2 = st.columns(2)
            with col_btn1:
                btn_disabled = (test_size == 1.0 or st.session_state.training_locked)
                if st.button("🚀 Mulai Training Model", disabled=btn_disabled, use_container_width=True, type="primary"):
                    with st.spinner("🔄 Melatih model..."):
                        progress_bar = st.progress(0)
                        for i in range(1, 101):
                            time.sleep(0.005)
                            progress_bar.progress(i)
                        res = run_training(df, test_size)
                        st.session_state.train_results = res
                        st.session_state.training_locked = True
                        st.session_state.batch_prediction_df = None
                    
                    if test_size == 0.0:
                        st.success("✅ Model berhasil dilatih! (Training Only)")
                    elif test_size < 1.0:
                        st.success("✅ Model berhasil dilatih!")
                    st.rerun()
            
            with col_btn2:
                if st.button("🔄 Reset Training", use_container_width=True):
                    st.session_state.train_results = None
                    st.session_state.training_locked = False
                    st.session_state.batch_prediction_df = None
                    st.rerun()
            
            if st.session_state.train_results:
                res = st.session_state.train_results
                st.markdown("---")
                st.markdown("### 📊 Ringkasan Hasil Training")
                
                c1, c2, c3, c4 = st.columns(4)
                with c1:
                    st.metric("Akurasi Training", f"{res['train_acc']*100:.2f}%" if res['train_acc'] is not None else "N/A")
                with c2:
                    if res['test_acc'] is not None:
                        st.metric("Akurasi Testing", f"{res['test_acc']*100:.2f}%")
                    else:
                        st.metric("Akurasi Testing", "N/A")
                with c3:
                    if res['y_pred'].size > 0:
                        benar = (res['y_test'].values == res['y_pred']).sum()
                        st.metric("Prediksi Benar", int(benar))
                    else:
                        st.metric("Prediksi Benar", "N/A")
                with c4:
                    if res['y_pred'].size > 0:
                        salah = (res['y_test'].values != res['y_pred']).sum()
                        st.metric("Prediksi Salah", int(salah))
                    else:
                        st.metric("Prediksi Salah", "N/A")
                
                if res['model'] is not None:
                    with st.expander("🧮 Parameter Model", expanded=False):
                        model = res['model']
                        fc = res['feature_cols']
                        st.dataframe(pd.DataFrame({
                            'Fitur': fc,
                            'Mean Tidak (0)': model.theta_[0].round(4),
                            'Mean Ya (1)': model.theta_[1].round(4),
                            'Std Tidak (0)': np.sqrt(model.var_[0]).round(4),
                            'Std Ya (1)': np.sqrt(model.var_[1]).round(4),
                        }), use_container_width=True, hide_index=True)

    elif page == "📈 Evaluasi":
        render_sticky_header("📈", "Evaluasi Model")
        
        if st.session_state.train_results is None:
            st.warning("⚠️ Model belum dilatih.")
        elif st.session_state.train_results['test_acc'] is None:
            st.warning("⚠️ Testing 0% - Gunakan upload data testing manual di Dashboard.")
        else:
            res = st.session_state.train_results
            th = make_plotly_theme()
            colors = make_plotly_colors()
            
            render_metric_cards([
                {"icon": "🏋️", "label": "Akurasi Training", "value": f"{res['train_acc']*100:.2f}%"},
                {"icon": "🧪", "label": "Akurasi Testing", "value": f"{res['test_acc']*100:.2f}%"},
                {"icon": "📊", "label": "F1-Score Macro", "value": f"{res['macro_f1']*100:.2f}%"},
                {"icon": "⚖️", "label": "F1-Score Weighted", "value": f"{res['w_f1']*100:.2f}%"},
            ])
            
            st.markdown("---")
            tabs = st.tabs(["📋 Tabel Metrik", "🗂️ Confusion Matrix", "📊 Grafik Evaluasi"])
            
            with tabs[0]:
                col1, col2 = st.columns(2)
                with col1:
                    st.markdown("**Per Kelas**")
                    st.dataframe(pd.DataFrame({
                        'Metrik': ['Precision', 'Recall', 'F1-Score'],
                        'Tidak Penerima (0)': [f"{res['prec0']:.4f}", f"{res['rec0']:.4f}", f"{res['f1_0']:.4f}"],
                        'Penerima (1)': [f"{res['prec1']:.4f}", f"{res['rec1']:.4f}", f"{res['f1_1']:.4f}"],
                    }), use_container_width=True, hide_index=True)
                with col2:
                    st.markdown("**Average**")
                    st.dataframe(pd.DataFrame({
                        'Tipe': ['Macro', 'Weighted'],
                        'Precision': [f"{res['macro_prec']:.4f}", f"{res['w_prec']:.4f}"],
                        'Recall': [f"{res['macro_rec']:.4f}", f"{res['w_rec']:.4f}"],
                        'F1-Score': [f"{res['macro_f1']:.4f}", f"{res['w_f1']:.4f}"],
                    }), use_container_width=True, hide_index=True)
            
            with tabs[1]:
                cm = res['cm']
                fig = go.Figure(data=go.Heatmap(
                    z=cm,
                    x=['Predicted: Tidak', 'Predicted: Ya'],
                    y=['Actual: Tidak', 'Actual: Ya'],
                    text=[[str(v) for v in row] for row in cm],
                    texttemplate="%{text}",
                    textfont={"size": 22, "family": "JetBrains Mono", "color": "#FFFFFF"},
                    colorscale=[[0, "#1E293B"], [0.5, "#3B82F6"], [1, "#93C5FD"]],
                    showscale=True, 
                    colorbar=dict(title="Jumlah", tickfont=dict(color="#FFFFFF"), title_font=dict(color="#FFFFFF"))
                ))
                fig.update_layout(**th, title="Confusion Matrix", height=420)
                st.plotly_chart(fig, use_container_width=True)
            
            with tabs[2]:
                metrics = {
                    'Accuracy Train': res['train_acc'], 'Accuracy Test': res['test_acc'],
                    'Prec (0)': res['prec0'], 'Recall (0)': res['rec0'], 'F1 (0)': res['f1_0'],
                    'Prec (1)': res['prec1'], 'Recall (1)': res['rec1'], 'F1 (1)': res['f1_1'],
                }
                fig1 = px.bar(
                    x=list(metrics.keys()), y=[v*100 for v in metrics.values()],
                    text=[f"{v*100:.1f}%" for v in metrics.values()],
                    color_discrete_sequence=[colors['primary']]*len(metrics),
                    title="Perbandingan Semua Metrik (%)"
                )
                fig1.update_traces(textfont_color='#FFFFFF')
                fig1.update_layout(**th, yaxis_title="Nilai (%)", height=400)
                st.plotly_chart(fig1, use_container_width=True)

    elif page == "🗂️ Hasil Prediksi":
        render_sticky_header("🗂️", "Hasil Prediksi")
        
        if st.session_state.train_results is None:
            st.warning("⚠️ Model belum dilatih.")
        elif st.session_state.train_results['y_pred'].size == 0:
            st.warning("⚠️ Testing 0% - Tidak ada hasil prediksi otomatis.")
        else:
            res = st.session_state.train_results
            df_r = res['results_df']
            
            benar = (df_r['Status'] == '✓ Benar').sum()
            salah = len(df_r) - benar
            
            render_metric_cards([
                {"icon": "📋", "label": "Total Data Uji", "value": str(len(df_r))},
                {"icon": "✓", "label": "Prediksi Benar", "value": str(benar), "desc": f"{benar/len(df_r)*100:.1f}%"},
                {"icon": "✗", "label": "Prediksi Salah", "value": str(salah), "desc": f"{salah/len(df_r)*100:.1f}%"},
                {"icon": "🎯", "label": "Akurasi", "value": f"{benar/len(df_r)*100:.2f}%"},
            ])
            
            st.markdown("---")
            col1, col2, col3 = st.columns(3)
            with col1:
                filter_status = st.selectbox("Filter Status", ["Semua", "✓ Benar", "✗ Salah"])
            with col2:
                filter_actual = st.selectbox("Filter Actual", ["Semua", "Penerima", "Tidak Penerima"])
            with col3:
                filter_pred = st.selectbox("Filter Predicted", ["Semua", "Penerima", "Tidak Penerima"])
            
            df_disp = df_r.copy()
            if filter_status != "Semua":
                df_disp = df_disp[df_disp['Status'] == filter_status]
            if filter_actual != "Semua":
                df_disp = df_disp[df_disp['Actual'] == filter_actual]
            if filter_pred != "Semua":
                df_disp = df_disp[df_disp['Predicted'] == filter_pred]
            
            st.caption(f"Menampilkan {len(df_disp)} dari {len(df_r)} prediksi")
            st.dataframe(df_disp, use_container_width=True, height=500, hide_index=True)

    elif page == "🔮 Prediksi Data Baru":
        show_prediction_page()

    elif page == "💾 Export Hasil":
        render_sticky_header("💾", "Export & Download")
        
        if st.session_state.df_processed is None:
            st.warning("⚠️ Dataset belum diproses.")
        else:
            df = st.session_state.df_processed
            info = st.session_state.preprocess_info
            le = info['le_pekerjaan']
            
            col1, col2 = st.columns(2)
            with col1:
                st.markdown("### 📊 Data Preprocessing")
                preproc_bytes = export_preprocessing_excel(df, le)
                st.download_button("⬇️ Download Data Preprocessing (.xlsx)", data=preproc_bytes,
                                  file_name="data_bsm_preprocessed.xlsx",
                                  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  use_container_width=True)
                
                st.markdown("---")
                st.markdown("### 📐 Data Normalisasi")
                norm_bytes = export_normalisasi_excel(df, info)
                st.download_button("⬇️ Download Hasil Normalisasi (.xlsx)", data=norm_bytes,
                                  file_name="data_normalisasi_zscore.xlsx",
                                  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  use_container_width=True)
            
            with col2:
                if st.session_state.train_results and st.session_state.train_results['y_pred'].size > 0:
                    res = st.session_state.train_results
                    st.markdown("### 🎯 Hasil Prediksi")
                    output = io.BytesIO()
                    with pd.ExcelWriter(output, engine='openpyxl') as writer:
                        res['results_df'].to_excel(writer, index=False, sheet_name='Hasil Prediksi')
                    output.seek(0)
                    st.download_button("⬇️ Download Hasil Prediksi (.xlsx)", data=output.getvalue(),
                                      file_name="hasil_prediksi_naive_bayes.xlsx",
                                      mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                      use_container_width=True)
                else:
                    st.info("ℹ️ Latih model terlebih dahulu untuk mengekspor hasil prediksi.")

    elif page == "ℹ️ Tentang Sistem":
        render_sticky_header("ℹ️", "Tentang Sistem")
        
        col1, col2, col3 = st.columns(3)
        with col1:
            render_metric_card("🎓", "Tujuan", "Klasifikasi BSM", "Membantu keputusan pemberian BSM secara objektif")
        with col2:
            render_metric_card("🔬", "Algoritma", "Gaussian NB", "Probabilistik berbasis Teorema Bayes")
        with col3:
            render_metric_card("📊", "Fitur", "5 Input", "Kelas, Pendapatan, Pekerjaan, Tanggungan, Status Rumah")
        
        st.markdown("---")
        st.markdown("### 🛠️ Teknologi")
        tech = [
            ("🐍", "Python 3.x", "Bahasa pemrograman"),
            ("⚡", "Streamlit", "Framework dashboard"),
            ("🐼", "Pandas", "Manipulasi data"),
            ("🤖", "Scikit-learn", "Machine Learning"),
            ("📈", "Plotly", "Visualisasi interaktif"),
            ("📗", "Openpyxl", "Ekspor Excel"),
        ]
        for icon, name, desc in tech:
            st.markdown(f"""
            <div style="display:flex; align-items:center; gap:14px; padding:12px 16px; 
                        background:#fff; border:3px solid #000; box-shadow:4px 4px 0 #000;
                        margin-bottom:8px;">
                <div style="font-size:1.6rem; width:44px; text-align:center;
                            background:#2146FF; border:2px solid #000; padding:6px;
                            box-shadow:2px 2px 0 #000; flex-shrink:0;">{icon}</div>
                <div>
                    <div style="font-family:'Archivo Black',sans-serif; font-weight:900; color:#000;
                                font-size:0.88rem; text-transform:uppercase;">{name}</div>
                    <div style="font-family:'IBM Plex Mono',monospace; font-size:0.75rem; color:#555;">{desc}</div>
                </div>
            </div>
            """, unsafe_allow_html=True)
    
    # Footer
    st.markdown("""
    <div class="footer">
        <div class="footer-brand">🎓 SISTEM KLASIFIKASI PENERIMA BSM</div>
        <div class="footer-divider"></div>
        <div style="font-family:'IBM Plex Mono',monospace;font-size:0.8rem;letter-spacing:0.06em;">
            GAUSSIAN NAIVE BAYES &nbsp;·&nbsp; MACHINE LEARNING &nbsp;·&nbsp; STREAMLIT
        </div>
        <div style="margin-top:0.5rem;font-family:'IBM Plex Mono',monospace;
                    color:rgba(255,255,255,0.4);font-size:0.7rem;letter-spacing:0.08em;">
            // DIBUAT UNTUK KEPERLUAN AKADEMIK DAN PENELITIAN //
        </div>
    </div>
    """, unsafe_allow_html=True)

# ─────────────────────────────────────────────────────────────
# APP ENTRY POINT
# ─────────────────────────────────────────────────────────────
if __name__ == "__main__":
    if st.session_state.app_state == 'home':
        show_landing()
    else:
        show_main()
